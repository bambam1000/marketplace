from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required
from django.contrib import messages as django_messages
from django.db.models import Sum, Count, F, Q, Avg
from django.db.models.functions import TruncMonth, TruncDay
from django.utils import timezone
from datetime import timedelta
from orders.models import Order, OrderItem
from catalog.models import Product, Category, Review
from accounts.models import User
from store.models import Store
import json
from messaging.models import Conversation, Message
from django.http import JsonResponse
from django.utils.text import slugify
from functools import wraps


def seller_or_admin_required(view_func):
    """Bloque l'accès aux acheteurs : réservé aux vendeurs (avec boutique) et admins."""
    @wraps(view_func)
    def wrapper(request, *args, **kwargs):
        user = request.user
        is_admin = user.is_superuser or user.role == 'admin'
        is_seller = user.is_seller and hasattr(user, 'store')
        is_member = user.store_memberships.filter(is_active=True).exists()
        if not (is_admin or is_seller or is_member):
            django_messages.error(request, "Accès réservé aux vendeurs et administrateurs.")
            return redirect('dashboard:index')
        return view_func(request, *args, **kwargs)
    return wrapper


def store_permission_required(permission):
    """Vérifie qu'un employé a la permission requise. Le propriétaire a tous les droits."""
    def decorator(view_func):
        @wraps(view_func)
        def wrapper(request, *args, **kwargs):
            user = request.user
            if user.is_superuser or user.role == 'admin' or hasattr(user, 'store'):
                return view_func(request, *args, **kwargs)
            member = user.store_memberships.filter(is_active=True).first()
            if member and member.has_permission(permission):
                return view_func(request, *args, **kwargs)
            django_messages.error(request, "Vous n'avez pas cette permission.")
            return redirect('dashboard:index')
        return wrapper
    return decorator


@login_required
def index(request):
    user = request.user
    is_admin = user.is_superuser or user.role == 'admin'
    is_seller = user.is_seller and hasattr(user, 'store')
    is_member = user.store_memberships.filter(is_active=True).exists()
    # Les acheteurs n'ont pas accès au dashboard (sauf employés de boutique)
    if not (is_admin or is_seller or is_member):
        django_messages.error(request, "Le tableau de bord est réservé aux vendeurs.")
        return redirect('home')
    now = timezone.now()
    today = now.date()
    store = getattr(user, 'store', None) if is_seller else None

    # Base querysets
    if is_admin:
        orders_qs = Order.objects.all()
        items_qs = OrderItem.objects.filter(order__is_paid=True)
        products_qs = Product.objects.all()
    elif is_seller:
        orders_qs = Order.objects.filter(items__store=store).distinct()
        items_qs = OrderItem.objects.filter(store=store, order__is_paid=True)
        products_qs = store.products.all()
    else:
        orders_qs = Order.objects.filter(buyer=user)
        items_qs = OrderItem.objects.none()
        products_qs = Product.objects.none()

    # KPIs
    total_revenue = items_qs.aggregate(t=Sum(F('price') * F('quantity')))['t'] or 0
    total_orders = orders_qs.count()
    pending = orders_qs.filter(status='pending').count()
    today_orders = orders_qs.filter(created_at__date=today).count()
    today_revenue = items_qs.filter(order__created_at__date=today).aggregate(
        t=Sum(F('price') * F('quantity')))['t'] or 0
    this_month_revenue = items_qs.filter(order__created_at__month=now.month, order__created_at__year=now.year).aggregate(
        t=Sum(F('price') * F('quantity')))['t'] or 0
    products_count = products_qs.count()
    low_stock = products_qs.filter(stock__lte=5, is_active=True).count()
    total_reviews = Review.objects.count() if is_admin else (Review.objects.filter(product__store=store).count() if is_seller else 0)
    avg_rating = Review.objects.aggregate(a=Avg('rating'))['a'] if is_admin else (Review.objects.filter(product__store=store).aggregate(a=Avg('rating'))['a'] if is_seller else 0)

    # Extra admin KPIs
    total_users = User.objects.count() if is_admin else 0
    total_sellers = User.objects.filter(role='seller').count() if is_admin else 0
    total_buyers = User.objects.filter(role='buyer').count() if is_admin else 0
    total_stores = Store.objects.count() if is_admin else 0
    verified_stores = Store.objects.filter(is_verified=True).count() if is_admin else 0
    new_users_30d = User.objects.filter(date_joined__gte=now - timedelta(days=30)).count() if is_admin else 0
    open_rfqs = 0
    unread_messages = 0
    if is_admin:
        from orders.models import RFQ
        from messaging.models import Message
        try:
            open_rfqs = RFQ.objects.filter(status='open').count()
            unread_messages = Message.objects.filter(is_read=False).count()
        except: pass

    # Seller-specific
    pending_payout = 0
    commission_total = 0
    if is_seller and store:
        from accounting.models import SellerWallet
        wallet = SellerWallet.objects.filter(user=user).first()
        if wallet:
            pending_payout = wallet.pending_balance
            commission_total = wallet.total_commission_paid

    # Charts
    monthly = items_qs.filter(
        order__created_at__gte=now - timedelta(days=180)
    ).annotate(month=TruncMonth('order__created_at')).values('month').annotate(
        rev=Sum(F('price') * F('quantity')), cnt=Count('order', distinct=True)
    ).order_by('month')
    months = [m['month'].strftime('%b') for m in monthly]
    revenues = [int(m['rev'] or 0) for m in monthly]
    order_counts = [m['cnt'] for m in monthly]

    # Daily (last 7 days)
    daily = items_qs.filter(
        order__created_at__gte=now - timedelta(days=7)
    ).annotate(day=TruncDay('order__created_at')).values('day').annotate(
        rev=Sum(F('price') * F('quantity'))
    ).order_by('day')
    daily_labels = [d['day'].strftime('%a') for d in daily]
    daily_data = [int(d['rev'] or 0) for d in daily]

    # Status distribution
    status_data = orders_qs.values('status').annotate(cnt=Count('id'))
    st_labels = [dict(Order.STATUS_CHOICES).get(s['status'], s['status']) for s in status_data]
    st_data = [s['cnt'] for s in status_data]

    # Top products
    top_products = items_qs.values('product__name').annotate(
        rev=Sum(F('price') * F('quantity')), qty=Sum('quantity')
    ).order_by('-rev')[:5]

    # Recent activity
    recent_orders = orders_qs.select_related('buyer').order_by('-created_at')[:8]

    return render(request, 'dashboard/index.html', {
        'is_admin': is_admin, 'is_seller': is_seller, 'store': store,
        'total_revenue': total_revenue, 'total_orders': total_orders,
        'pending_orders': pending, 'products_count': products_count,
        'today_orders': today_orders, 'today_revenue': today_revenue,
        'this_month_revenue': this_month_revenue,
        'low_stock': low_stock, 'total_reviews': total_reviews,
        'avg_rating': round(avg_rating or 0, 1),
        'total_users': total_users, 'total_sellers': total_sellers,
        'total_buyers': total_buyers, 'total_stores': total_stores,
        'verified_stores': verified_stores, 'new_users_30d': new_users_30d,
        'open_rfqs': open_rfqs, 'unread_messages': unread_messages,
        'pending_payout': pending_payout, 'commission_total': commission_total,
        'months_json': json.dumps(months), 'revenues_json': json.dumps(revenues),
        'order_counts_json': json.dumps(order_counts),
        'daily_labels_json': json.dumps(daily_labels), 'daily_data_json': json.dumps(daily_data),
        'st_labels_json': json.dumps(st_labels), 'st_data_json': json.dumps(st_data),
        'top_products': top_products, 'recent_orders': recent_orders,
    })


@login_required
@seller_or_admin_required
def dash_orders(request):
    if request.user.is_seller and hasattr(request.user, 'store'):
        orders = Order.objects.filter(items__store=request.user.store).distinct()
    else:
        orders = Order.objects.all()
    status = request.GET.get('status')
    if status: orders = orders.filter(status=status)
    q = request.GET.get('q', '').strip()
    if q:
        orders = orders.filter(Q(order_number__icontains=q) | Q(buyer__username__icontains=q) | Q(buyer__first_name__icontains=q) | Q(buyer__last_name__icontains=q))

    # Filtres date
    from datetime import datetime
    date_from = request.GET.get('from', '')
    date_to = request.GET.get('to', '')
    if date_from:
        try:
            orders = orders.filter(created_at__date__gte=datetime.strptime(date_from, '%Y-%m-%d').date())
        except ValueError:
            pass
    if date_to:
        try:
            orders = orders.filter(created_at__date__lte=datetime.strptime(date_to, '%Y-%m-%d').date())
        except ValueError:
            pass

    total_revenue = OrderItem.objects.filter(order__in=orders, order__is_paid=True).aggregate(t=Sum(F('price') * F('quantity')))['t'] or 0
    return render(request, 'dashboard/orders.html', {
        'orders': orders.select_related('buyer').order_by('-created_at'),
        'status_choices': Order.STATUS_CHOICES,
        'current_status': status,
        'search_query': q,
        'date_from': date_from,
        'date_to': date_to,
        'total': orders.count(),
        'pending': orders.filter(status='pending').count(),
        'delivered': orders.filter(status='delivered').count(),
        'total_revenue': total_revenue,
    })


@login_required
@seller_or_admin_required
def dash_order_detail(request, order_number):
    order = get_object_or_404(Order, order_number=order_number)
    # Un vendeur ne peut gérer que les commandes contenant ses produits
    is_admin = request.user.is_superuser or request.user.role == 'admin'
    if not is_admin:
        if not order.items.filter(store=request.user.store).exists():
            django_messages.error(request, "Cette commande ne concerne pas votre boutique.")
            return redirect('dashboard:orders')
    if request.method == 'POST':
        new_status = request.POST.get('status')
        if new_status:
            order.status = new_status
            if new_status in ['confirmed', 'delivered']:
                order.is_paid = True
                # Create accounting transactions
                from accounting.models import Transaction, SellerWallet
                for item in order.items.all():
                    sale_amount = int(item.price) * item.quantity
                    commission = int(sale_amount * 0.10)
                    Transaction.objects.get_or_create(
                        order=order, user=item.store.owner, type='sale',
                        defaults={
                            'amount': sale_amount,
                            'commission_amount': commission,
                            'net_amount': sale_amount - commission,
                            'status': 'completed',
                            'reference': f'Vente #{order.order_number}',
                        }
                    )
                    wallet, _ = SellerWallet.objects.get_or_create(user=item.store.owner)
                    wallet.balance += (sale_amount - commission)
                    wallet.total_earned += sale_amount
                    wallet.total_commission_paid += commission
                    wallet.save()
            order.save()
            # Notifier le client du changement de statut
            from messaging.utils import notify
            status_labels = {
                'confirmed': 'confirmée', 'processing': 'en préparation',
                'shipped': 'expédiée', 'delivered': 'livrée', 'cancelled': 'annulée',
            }
            label = status_labels.get(new_status)
            if label:
                notify(
                    order.buyer, 'order',
                    f'Commande {order.order_number} {label}',
                    f'Votre commande de {order.total_amount:.0f} FCFA est maintenant {label}.',
                    url=f'/commandes/{order.order_number}/',
                    send_email=True,
                )
            django_messages.success(request, f'Statut mis à jour.')
        tracking = request.POST.get('tracking_number')
        if tracking:
            order.tracking_number = tracking
            order.save()
        return redirect('dashboard:order_detail', order_number=order_number)
    return render(request, 'dashboard/order_detail.html', {'order': order})


def get_user_warehouse(request):
    """Retourne l'entrepôt de l'employé connecté, ou None pour le propriétaire/admin."""
    member = request.user.store_memberships.filter(is_active=True).first()
    return member.warehouse if member and member.warehouse else None


@login_required
@seller_or_admin_required
def dash_products(request):
    if request.user.is_seller and hasattr(request.user, 'store'):
        products = request.user.store.products.all()
    else:
        products = Product.objects.all()

    # Filtre par entrepôt (paramètre ou entrepôt de l'employé)
    from inventory.models import Warehouse
    warehouse_id = request.GET.get('warehouse')
    member_warehouse = get_user_warehouse(request)
    if member_warehouse:
        warehouse_id = member_warehouse.id
    current_warehouse = None
    if warehouse_id:
        current_warehouse = Warehouse.objects.filter(pk=warehouse_id).first()
        if current_warehouse:
            products = products.filter(warehouse_stocks__warehouse=current_warehouse).distinct()

    q = request.GET.get('q', '').strip()
    if q:
        products = products.filter(name__icontains=q)
    stock_filter = request.GET.get('stock', '')
    if stock_filter == 'out':
        products = products.filter(stock=0)
    elif stock_filter == 'low':
        products = [p for p in products if p.is_low_stock]
    elif stock_filter == 'archived':
        products = products.filter(is_active=False).select_related('category', 'store')
    else:
        products = products.select_related('category', 'store')
    return render(request, 'dashboard/products.html', {
        'products': products,
        'categories': Category.objects.filter(is_active=True),
        'search_query': q,
        'stock_filter': stock_filter,
        'warehouse': current_warehouse,
        'current_warehouse': current_warehouse,
        'total': len(products),
        'active': sum(1 for p in products if p.is_active),
        'low': sum(1 for p in products if p.is_low_stock),
        'featured': sum(1 for p in products if p.is_featured),
    })


@login_required
@seller_or_admin_required
def dash_product_edit(request, pk=None):
    product = get_object_or_404(Product, pk=pk) if pk else None
    is_admin = request.user.is_superuser or request.user.role == 'admin'
    # Un vendeur ne peut modifier que ses propres produits
    if product and not is_admin and product.store != request.user.store:
        django_messages.error(request, "Vous ne pouvez pas modifier ce produit.")
        return redirect('dashboard:products')
    categories = Category.objects.all()
    store = getattr(request.user, 'store', None) or (Store.objects.first() if is_admin else None)
    if store is None:
        django_messages.error(request, "Aucune boutique associée.")
        return redirect('dashboard:index')
    if request.method == 'POST':
        data = {
            'name': request.POST.get('name'),
            'description': request.POST.get('description'),
            'price': request.POST.get('price'),
            'old_price': request.POST.get('old_price') or None,
            'stock': request.POST.get('stock', 0),
            'min_order': request.POST.get('min_order', 1),
            'colors': request.POST.get('colors', ''),
            'sizes': request.POST.get('sizes', ''),
            'specifications': request.POST.get('specifications', ''),
            'origin': request.POST.get('origin', 'Cameroun'),
            'is_active': 'is_active' in request.POST,
            'is_featured': 'is_featured' in request.POST,
            'is_flash_deal': 'is_flash_deal' in request.POST,
        }
        cat_id = request.POST.get('category')
        if cat_id: data['category'] = Category.objects.get(pk=cat_id)
        if product:
            # Le stock ne se modifie que via l'ajustement (traçabilité)
            data.pop('stock', None)
            for k, v in data.items(): setattr(product, k, v)
            for f in ['image', 'image_2', 'image_3', 'image_4']:
                if request.FILES.get(f): setattr(product, f, request.FILES[f])
            product.save()
        else:
            data['store'] = store
            product = Product(**data)
            for f in ['image', 'image_2', 'image_3', 'image_4']:
                if request.FILES.get(f): setattr(product, f, request.FILES[f])
            product.save()
        django_messages.success(request, 'Produit sauvegardé !')
        return redirect('dashboard:products')
    return render(request, 'dashboard/product_edit.html', {'product': product, 'categories': categories})


@login_required
@seller_or_admin_required
def dash_product_delete(request, pk):
    """Archive un produit (style Odoo : on ne supprime pas, on archive).
    La suppression définitive n'est possible que si le produit n'a aucun historique."""
    product = get_object_or_404(Product, pk=pk)
    is_admin = request.user.is_superuser or request.user.role == 'admin'
    if not is_admin and product.store != request.user.store:
        django_messages.error(request, 'Vous ne pouvez pas modifier ce produit.')
        return redirect('dashboard:products')
    if request.method == 'POST':
        has_history = product.stock_movements.exists() or product.orderitem_set.exists()
        if has_history:
            # Archivage : le produit reste en base pour la traçabilité
            product.is_active = False
            product.save(update_fields=['is_active'])
            django_messages.success(request, f'Produit "{product.name}" archivé. Il n\'est plus visible en boutique mais son historique est conservé.')
        else:
            name = product.name
            product.delete()
            django_messages.success(request, f'Produit "{name}" supprimé (aucun historique).')
    return redirect('dashboard:products')


@login_required
@seller_or_admin_required
def dash_product_unarchive(request, pk):
    """Réactive un produit archivé."""
    product = get_object_or_404(Product, pk=pk)
    is_admin = request.user.is_superuser or request.user.role == 'admin'
    if not is_admin and product.store != request.user.store:
        django_messages.error(request, 'Vous ne pouvez pas modifier ce produit.')
        return redirect('dashboard:products')
    if request.method == 'POST':
        product.is_active = True
        product.save(update_fields=['is_active'])
        django_messages.success(request, f'Produit "{product.name}" réactivé.')
    return redirect('dashboard:products')


@login_required
@seller_or_admin_required
def dash_sales(request):
    """Page Ventes : commandes payées avec statistiques"""
    if request.user.is_seller and hasattr(request.user, 'store'):
        items = OrderItem.objects.filter(store=request.user.store, order__is_paid=True)
    else:
        items = OrderItem.objects.filter(order__is_paid=True)

    # Filtre par entrepôt
    from inventory.models import Warehouse
    warehouse_id = request.GET.get('warehouse')
    member_warehouse = get_user_warehouse(request)
    if member_warehouse:
        warehouse_id = member_warehouse.id
    current_warehouse = None
    if warehouse_id:
        items = items.filter(warehouse_id=warehouse_id)
        current_warehouse = Warehouse.objects.filter(pk=warehouse_id).first()

    q = request.GET.get('q', '').strip()
    if q:
        items = items.filter(Q(product__name__icontains=q) | Q(order__order_number__icontains=q))

    now = timezone.now()
    total_revenue = items.aggregate(t=Sum(F('price') * F('quantity')))['t'] or 0
    total_qty = items.aggregate(t=Sum('quantity'))['t'] or 0
    month_revenue = items.filter(order__created_at__month=now.month, order__created_at__year=now.year).aggregate(t=Sum(F('price') * F('quantity')))['t'] or 0
    orders_count = items.values('order').distinct().count()
    avg_basket = int(total_revenue / orders_count) if orders_count else 0

    store = getattr(request.user, 'store', None)
    return render(request, 'dashboard/sales.html', {
        'items': items.select_related('order', 'order__buyer', 'product').order_by('-order__created_at')[:100],
        'total_revenue': total_revenue,
        'total_qty': total_qty,
        'month_revenue': month_revenue,
        'orders_count': orders_count,
        'avg_basket': avg_basket,
        'search_query': q,
        'my_products': store.products.filter(is_active=True, stock__gt=0) if store else [],
        'warehouse': current_warehouse,
        'payment_choices': Order.PAYMENT_CHOICES,
    })


@login_required
@seller_or_admin_required
def dash_sale_create(request):
    """Vente directe : vendre sans commande en ligne (comptoir, téléphone...)"""
    if request.method == 'POST':
        store = getattr(request.user, 'store', None)
        if not store:
            django_messages.error(request, "Vous devez avoir une boutique.")
            return redirect('dashboard:sales')

        product_ids = request.POST.getlist('product[]')
        quantities = request.POST.getlist('quantity[]')
        customer_name = request.POST.get('customer_name', '').strip() or 'Client comptoir'
        customer_phone = request.POST.get('customer_phone', '').strip()
        payment_method = request.POST.get('payment_method', 'cash')

        # Valider les lignes
        lines = []
        for i, pid in enumerate(product_ids):
            if not pid:
                continue
            product = get_object_or_404(Product, pk=pid, store=store)
            qty = int(quantities[i]) if i < len(quantities) else 1
            if qty < 1:
                continue
            if qty > product.stock:
                django_messages.error(request, f'Stock insuffisant pour "{product.name}" : {product.stock} disponible(s).')
                return redirect('dashboard:sales')
            lines.append((product, qty))

        if not lines:
            django_messages.error(request, 'Ajoutez au moins un produit.')
            return redirect('dashboard:sales')

        from inventory.models import Warehouse
        warehouse = Warehouse.objects.filter(store=store, is_default=True).first()

        total = sum(p.price * q for p, q in lines)
        order = Order.objects.create(
            buyer=request.user,
            status='delivered',
            is_paid=True,
            payment_method=payment_method,
            subtotal=total,
            shipping_cost=0,
            total_amount=total,
            shipping_name=customer_name,
            shipping_phone=customer_phone or '—',
            shipping_address='Vente directe',
            shipping_city=store.city,
            notes=f'Vente directe — {customer_name}',
        )
        for product, qty in lines:
            OrderItem.objects.create(order=order, product=product, store=store, warehouse=warehouse, quantity=qty, price=product.price)
            product.orders_count += qty
            product.save(update_fields=['orders_count'])
            product.adjust_stock(-qty, 'sale', user=request.user, reason='Vente directe', reference=order.order_number, warehouse=warehouse)

        django_messages.success(request, f'Vente {order.order_number} enregistrée : {len(lines)} produit(s), {total:,} F.')
    return redirect('dashboard:sales')


# ======== DEVIS RFQ (dashboard vendeur) ========
@login_required
@seller_or_admin_required
def dash_rfqs(request):
    """Liste des demandes de devis ouvertes + création de demande — dans le dashboard"""
    from orders.models import RFQ

    # Le vendeur peut aussi créer une demande de devis (il achète aussi)
    if request.method == 'POST':
        RFQ.objects.create(
            buyer=request.user,
            product_name=request.POST.get('product_name'),
            category_id=request.POST.get('category') or None,
            description=request.POST.get('description'),
            quantity=int(request.POST.get('quantity', 1)),
            target_price=request.POST.get('target_price') or None,
            unit=request.POST.get('unit', 'pièce'),
        )
        django_messages.success(request, 'Votre demande de devis a été publiée !')
        return redirect('dashboard:rfqs')

    from orders.models import Quote
    view = request.GET.get('view', 'open')

    # RFQ où le vendeur a déjà soumis un devis
    my_quoted_ids = []
    if hasattr(request.user, 'store'):
        my_quoted_ids = list(Quote.objects.filter(seller=request.user).values_list('rfq_id', flat=True))

    # Mes propres demandes de devis
    my_rfqs = RFQ.objects.filter(buyer=request.user).annotate(quote_cnt=Count('quotes')).order_by('-created_at')

    # Liste principale selon l'onglet
    if view == 'mine':
        rfqs = my_rfqs
    elif view == 'quoted':
        rfqs = RFQ.objects.filter(pk__in=my_quoted_ids).annotate(quote_cnt=Count('quotes')).order_by('-created_at')
    else:
        rfqs = RFQ.objects.filter(status='open').exclude(buyer=request.user).annotate(quote_cnt=Count('quotes')).order_by('-created_at')

    q = request.GET.get('q', '').strip()
    if q:
        rfqs = rfqs.filter(Q(product_name__icontains=q) | Q(description__icontains=q))
    category = request.GET.get('category')
    if category:
        rfqs = rfqs.filter(category_id=category)

    return render(request, 'dashboard/rfqs.html', {
        'rfqs': rfqs,
        'my_rfqs': my_rfqs,
        'categories': Category.objects.filter(parent__isnull=True),
        'search_query': q,
        'category_filter': category,
        'current_view': view,
        'my_quoted_ids': my_quoted_ids,
        'total_open': RFQ.objects.filter(status='open').exclude(buyer=request.user).count(),
        'my_quotes_count': len(my_quoted_ids),
        'my_rfqs_count': my_rfqs.count(),
    })


@login_required
@seller_or_admin_required
def dash_rfq_detail(request, rfq_id):
    """Détail RFQ + soumission de devis — dans le dashboard"""
    from orders.models import RFQ, Quote
    rfq = get_object_or_404(RFQ, pk=rfq_id)
    user_quote = Quote.objects.filter(rfq=rfq, seller=request.user).first()

    if request.method == 'POST':
        if not hasattr(request.user, 'store'):
            django_messages.error(request, 'Vous devez avoir une boutique.')
            return redirect('dashboard:rfq_detail', rfq_id=rfq_id)
        if rfq.status != 'open':
            django_messages.error(request, "Cette demande n'est plus ouverte.")
            return redirect('dashboard:rfq_detail', rfq_id=rfq_id)

        price_per_unit = int(request.POST.get('price_per_unit'))
        Quote.objects.update_or_create(
            rfq=rfq, seller=request.user,
            defaults={
                'store': request.user.store,
                'price_per_unit': price_per_unit,
                'total_price': price_per_unit * rfq.quantity,
                'delivery_time': request.POST.get('delivery_time') or '',
                'payment_terms': request.POST.get('payment_terms') or '',
                'description': request.POST.get('description') or '',
            }
        )
        if rfq.status == 'open':
            rfq.status = 'quoted'
            rfq.save()
        django_messages.success(request, 'Votre devis a été soumis avec succès !')
        return redirect('dashboard:rfq_detail', rfq_id=rfq_id)

    return render(request, 'dashboard/rfq_detail.html', {
        'rfq': rfq,
        'user_quote': user_quote,
    })


# ======== ENTREPÔTS ========
@login_required
@seller_or_admin_required
def dash_warehouses(request):
    """Liste des entrepôts + création"""
    from inventory.models import Warehouse
    store = getattr(request.user, 'store', None)
    is_admin = request.user.is_superuser or request.user.role == 'admin'

    if request.method == 'POST':
        if not store:
            django_messages.error(request, "Vous devez avoir une boutique.")
            return redirect('dashboard:warehouses')
        wh = Warehouse.objects.create(
            store=store,
            name=request.POST.get('name'),
            code=request.POST.get('code', '').upper(),
            address=request.POST.get('address', ''),
            city=request.POST.get('city', store.city),
        )
        # Lier à des boutiques supplémentaires du même propriétaire
        linked = request.POST.getlist('linked_stores')
        if linked:
            wh.linked_stores.set(Store.objects.filter(pk__in=linked, owner=request.user))
        django_messages.success(request, 'Entrepôt créé !')
        return redirect('dashboard:warehouses')

    warehouses = Warehouse.objects.all() if is_admin else (store.warehouses.all() if store else Warehouse.objects.none())
    other_stores = Store.objects.filter(owner=request.user).exclude(pk=store.pk) if store else []
    return render(request, 'dashboard/warehouses.html', {
        'warehouses': warehouses,
        'other_stores': other_stores,
    })


@login_required
@seller_or_admin_required
def dash_warehouse_detail(request, pk):
    """Détail d'un entrepôt : stocks par produit"""
    from inventory.models import Warehouse
    warehouse = get_object_or_404(Warehouse, pk=pk)
    is_admin = request.user.is_superuser or request.user.role == 'admin'
    if not is_admin and warehouse.store != request.user.store:
        django_messages.error(request, "Accès refusé.")
        return redirect('dashboard:warehouses')

    stocks = warehouse.stocks.select_related('product').order_by('product__name')
    q = request.GET.get('q', '').strip()
    if q:
        stocks = stocks.filter(product__name__icontains=q)

    return render(request, 'dashboard/warehouse_detail.html', {
        'warehouse': warehouse,
        'stocks': stocks,
        'search_query': q,
        'products': warehouse.store.products.filter(is_active=True),
        'other_warehouses': warehouse.store.warehouses.filter(is_active=True).exclude(pk=warehouse.pk),
    })


@login_required
@seller_or_admin_required
def dash_warehouse_reception(request, pk):
    """Réception de marchandises dans un entrepôt (entrée de stock)"""
    from inventory.models import Warehouse
    warehouse = get_object_or_404(Warehouse, pk=pk)
    is_admin = request.user.is_superuser or request.user.role == 'admin'
    if not is_admin and warehouse.store != request.user.store:
        django_messages.error(request, "Accès refusé.")
        return redirect('dashboard:warehouses')

    if request.method == 'POST':
        product_ids = request.POST.getlist('product[]')
        quantities = request.POST.getlist('quantity[]')
        reason = request.POST.get('reason', '').strip() or 'Réception marchandises'
        count = 0
        for i, pid in enumerate(product_ids):
            if not pid:
                continue
            product = get_object_or_404(Product, pk=pid, store=warehouse.store)
            qty = int(quantities[i]) if i < len(quantities) else 0
            if qty > 0:
                product.adjust_stock(qty, 'in', user=request.user, reason=reason,
                                     reference=warehouse.code, warehouse=warehouse)
                count += 1
        if count:
            django_messages.success(request, f'Réception enregistrée : {count} produit(s) ajouté(s) à {warehouse.name}.')
        else:
            django_messages.error(request, 'Aucun produit valide.')
    return redirect('dashboard:warehouse_detail', pk=pk)


@login_required
@seller_or_admin_required
def dash_warehouse_employees(request, pk):
    """Employés assignés à un entrepôt précis"""
    from inventory.models import Warehouse
    warehouse = get_object_or_404(Warehouse, pk=pk)
    is_admin = request.user.is_superuser or request.user.role == 'admin'
    if not is_admin and warehouse.store != request.user.store:
        django_messages.error(request, "Accès refusé.")
        return redirect('dashboard:warehouses')

    members = warehouse.store.members.filter(warehouse=warehouse).select_related('user', 'role')
    return render(request, 'dashboard/warehouse_employees.html', {
        'warehouse': warehouse,
        'members': members,
    })


@login_required
@seller_or_admin_required
def dash_warehouse_orders(request, pk):
    """Commandes contenant des produits d'un entrepôt"""
    from inventory.models import Warehouse
    warehouse = get_object_or_404(Warehouse, pk=pk)
    is_admin = request.user.is_superuser or request.user.role == 'admin'
    if not is_admin and warehouse.store != request.user.store:
        django_messages.error(request, "Accès refusé.")
        return redirect('dashboard:warehouses')

    orders = Order.objects.filter(items__warehouse=warehouse).distinct().select_related('buyer').order_by('-created_at')
    status = request.GET.get('status', '')
    if status:
        orders = orders.filter(status=status)
    q = request.GET.get('q', '').strip()
    if q:
        orders = orders.filter(Q(order_number__icontains=q) | Q(buyer__username__icontains=q))

    return render(request, 'dashboard/warehouse_orders.html', {
        'warehouse': warehouse,
        'orders': orders,
        'status_choices': Order.STATUS_CHOICES,
        'current_status': status,
        'search_query': q,
    })


@login_required
@seller_or_admin_required
def dash_warehouse_stores(request, pk):
    """Boutiques alimentées par un entrepôt"""
    from inventory.models import Warehouse
    warehouse = get_object_or_404(Warehouse, pk=pk)
    is_admin = request.user.is_superuser or request.user.role == 'admin'
    if not is_admin and warehouse.store != request.user.store:
        django_messages.error(request, "Accès refusé.")
        return redirect('dashboard:warehouses')
    return render(request, 'dashboard/warehouse_stores.html', {'warehouse': warehouse})


@login_required
@seller_or_admin_required
def dash_warehouse_stock(request, pk):
    """Stock détaillé d'un entrepôt"""
    from inventory.models import Warehouse
    warehouse = get_object_or_404(Warehouse, pk=pk)
    is_admin = request.user.is_superuser or request.user.role == 'admin'
    if not is_admin and warehouse.store != request.user.store:
        django_messages.error(request, "Accès refusé.")
        return redirect('dashboard:warehouses')

    stocks = warehouse.stocks.select_related('product').order_by('product__name')
    q = request.GET.get('q', '').strip()
    if q:
        stocks = stocks.filter(product__name__icontains=q)
    return render(request, 'dashboard/warehouse_stock.html', {
        'warehouse': warehouse,
        'stocks': stocks,
        'search_query': q,
    })


@login_required
@seller_or_admin_required
def dash_warehouse_accounting(request, pk):
    """Comptabilité d'un entrepôt : revenus, dépenses, résultat"""
    from inventory.models import Warehouse
    from accounting.models import WarehouseExpense
    from datetime import datetime

    warehouse = get_object_or_404(Warehouse, pk=pk)
    is_admin = request.user.is_superuser or request.user.role == 'admin'
    if not is_admin and warehouse.store != request.user.store:
        django_messages.error(request, "Accès refusé.")
        return redirect('dashboard:warehouses')

    # Ajouter une dépense
    if request.method == 'POST':
        WarehouseExpense.objects.create(
            warehouse=warehouse,
            category=request.POST.get('category', 'other'),
            description=request.POST.get('description'),
            amount=int(request.POST.get('amount', 0)),
            date=request.POST.get('date') or timezone.now().date(),
            created_by=request.user,
        )
        django_messages.success(request, 'Dépense enregistrée.')
        return redirect('dashboard:warehouse_accounting', pk=pk)

    # Période (mois en cours par défaut)
    now = timezone.now()
    month = int(request.GET.get('month', now.month))
    year = int(request.GET.get('year', now.year))

    # Revenus de l'entrepôt (ventes payées)
    revenue = OrderItem.objects.filter(
        warehouse=warehouse, order__is_paid=True,
        order__created_at__month=month, order__created_at__year=year,
    ).aggregate(t=Sum(F('price') * F('quantity')))['t'] or 0

    # Commissions (10%)
    from decimal import Decimal
    commission = int(revenue * Decimal('0.10'))

    # Dépenses de l'entrepôt
    expenses = warehouse.expenses.filter(date__month=month, date__year=year)
    total_expenses = expenses.aggregate(t=Sum('amount'))['t'] or 0

    # Résultat net
    net_profit = revenue - commission - total_expenses

    # Revenus par mois (6 derniers mois) pour le graphique
    monthly = []
    for i in range(5, -1, -1):
        d = now - timedelta(days=30 * i)
        rev = OrderItem.objects.filter(
            warehouse=warehouse, order__is_paid=True,
            order__created_at__month=d.month, order__created_at__year=d.year,
        ).aggregate(t=Sum(F('price') * F('quantity')))['t'] or 0
        monthly.append({'label': d.strftime('%b'), 'revenue': int(rev)})

    return render(request, 'dashboard/warehouse_accounting.html', {
        'warehouse': warehouse,
        'revenue': revenue,
        'commission': commission,
        'total_expenses': total_expenses,
        'net_profit': net_profit,
        'expenses': expenses,
        'monthly': monthly,
        'current_month': month,
        'current_year': year,
        'months': [(i, timezone.datetime(2000, i, 1).strftime('%B')) for i in range(1, 13)],
        'years': range(now.year - 2, now.year + 1),
        'expense_categories': WarehouseExpense.CATEGORY_CHOICES,
    })


@login_required
@seller_or_admin_required
def dash_warehouse_expenses(request, pk):
    """Liste des dépenses d'un entrepôt avec filtres"""
    from inventory.models import Warehouse
    from accounting.models import WarehouseExpense
    from datetime import datetime

    warehouse = get_object_or_404(Warehouse, pk=pk)
    is_admin = request.user.is_superuser or request.user.role == 'admin'
    if not is_admin and warehouse.store != request.user.store:
        django_messages.error(request, "Accès refusé.")
        return redirect('dashboard:warehouses')

    expenses = warehouse.expenses.all()

    # Filtres
    category = request.GET.get('category', '')
    if category:
        expenses = expenses.filter(category=category)
    date_from = request.GET.get('from', '')
    date_to = request.GET.get('to', '')
    if date_from:
        try:
            expenses = expenses.filter(date__gte=datetime.strptime(date_from, '%Y-%m-%d').date())
        except ValueError:
            pass
    if date_to:
        try:
            expenses = expenses.filter(date__lte=datetime.strptime(date_to, '%Y-%m-%d').date())
        except ValueError:
            pass

    total = expenses.aggregate(t=Sum('amount'))['t'] or 0

    return render(request, 'dashboard/warehouse_expenses.html', {
        'warehouse': warehouse,
        'expenses': expenses,
        'total': total,
        'category_filter': category,
        'date_from': date_from,
        'date_to': date_to,
        'categories': WarehouseExpense.CATEGORY_CHOICES,
        'today': timezone.now().date().isoformat(),
    })


# ======== TRANSFERTS ========
@login_required
@seller_or_admin_required
def dash_transfers(request):
    """Liste des transferts + création"""
    from inventory.models import StockTransfer, TransferItem, Warehouse
    store = getattr(request.user, 'store', None)
    is_admin = request.user.is_superuser or request.user.role == 'admin'

    if request.method == 'POST':
        if not store:
            django_messages.error(request, "Vous devez avoir une boutique.")
            return redirect('dashboard:transfers')
        from_wh = get_object_or_404(Warehouse, pk=request.POST.get('from_warehouse'), store=store)
        to_wh = get_object_or_404(Warehouse, pk=request.POST.get('to_warehouse'), store=store)
        if from_wh == to_wh:
            django_messages.error(request, "Les deux entrepôts doivent être différents.")
            return redirect('dashboard:transfers')

        transfer = StockTransfer.objects.create(
            store=store, from_warehouse=from_wh, to_warehouse=to_wh,
            transfer_type=request.POST.get('transfer_type', 'transfer'),
            notes=request.POST.get('notes', ''), created_by=request.user,
        )
        product_ids = request.POST.getlist('product[]')
        quantities = request.POST.getlist('quantity[]')
        for i, pid in enumerate(product_ids):
            if pid:
                qty = int(quantities[i]) if i < len(quantities) else 1
                if qty > 0:
                    TransferItem.objects.create(transfer=transfer, product_id=pid, quantity=qty)
        django_messages.success(request, f'Transfert {transfer.reference} créé en brouillon.')
        return redirect('dashboard:transfers')

    transfers = StockTransfer.objects.all() if is_admin else (store.transfers.all() if store else StockTransfer.objects.none())
    status = request.GET.get('status', '')
    if status:
        transfers = transfers.filter(status=status)
    ttype = request.GET.get('ttype', '')
    if ttype:
        transfers = transfers.filter(transfer_type=ttype)

    # Recherche par référence
    q = request.GET.get('q', '').strip()
    if q:
        transfers = transfers.filter(Q(reference__icontains=q) | Q(from_warehouse__name__icontains=q) | Q(to_warehouse__name__icontains=q))

    # Filtres date
    from datetime import datetime
    date_from = request.GET.get('from', '')
    date_to = request.GET.get('to', '')
    if date_from:
        try:
            transfers = transfers.filter(created_at__date__gte=datetime.strptime(date_from, '%Y-%m-%d').date())
        except ValueError:
            pass
    if date_to:
        try:
            transfers = transfers.filter(created_at__date__lte=datetime.strptime(date_to, '%Y-%m-%d').date())
        except ValueError:
            pass

    return render(request, 'dashboard/transfers.html', {
        'transfers': transfers.select_related('from_warehouse', 'to_warehouse'),
        'warehouses': store.warehouses.filter(is_active=True) if store else [],
        'products': store.products.filter(is_active=True) if store else [],
        'status_filter': status,
        'status_choices': StockTransfer.STATUS_CHOICES,
        'type_filter': ttype,
        'search_query': q,
        'date_from': date_from,
        'date_to': date_to,
    })


def _filtered_transfers(request):
    """Transferts filtrés selon les paramètres GET."""
    from inventory.models import StockTransfer
    from datetime import datetime
    store = getattr(request.user, 'store', None)
    is_admin = request.user.is_superuser or request.user.role == 'admin'
    transfers = StockTransfer.objects.all() if is_admin else (store.transfers.all() if store else StockTransfer.objects.none())
    status = request.GET.get('status', '')
    if status:
        transfers = transfers.filter(status=status)
    date_from = request.GET.get('from', '')
    date_to = request.GET.get('to', '')
    if date_from:
        try:
            transfers = transfers.filter(created_at__date__gte=datetime.strptime(date_from, '%Y-%m-%d').date())
        except ValueError:
            pass
    if date_to:
        try:
            transfers = transfers.filter(created_at__date__lte=datetime.strptime(date_to, '%Y-%m-%d').date())
        except ValueError:
            pass
    return transfers.select_related('from_warehouse', 'to_warehouse', 'created_by')


@login_required
@seller_or_admin_required
def dash_transfers_export_pdf(request):
    """Export PDF des transferts filtrés"""
    from django.http import HttpResponse
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas
    import io

    transfers = _filtered_transfers(request)[:500]
    buffer = io.BytesIO()
    p = canvas.Canvas(buffer, pagesize=landscape(A4))
    width, height = landscape(A4)

    p.setFont('Helvetica-Bold', 16)
    p.drawString(15*mm, height - 18*mm, 'AfriMarket — Mouvements entre entrepôts')
    p.setFont('Helvetica', 9)
    p.drawString(15*mm, height - 24*mm, f'Généré le {timezone.now().strftime("%d/%m/%Y à %H:%M")} — {transfers.count()} transfert(s)')
    p.line(15*mm, height - 27*mm, width - 15*mm, height - 27*mm)

    y = height - 36*mm
    p.setFont('Helvetica-Bold', 8)
    headers = ['Référence', 'De', 'Vers', 'Produits', 'Statut', 'Créé le', 'Par']
    cols = [15, 50, 95, 140, 165, 200, 235]
    for x, h in zip(cols, headers):
        p.drawString(x*mm, y, h)
    p.line(15*mm, y - 2*mm, width - 15*mm, y - 2*mm)
    y -= 8*mm

    p.setFont('Helvetica', 8)
    for t in transfers:
        if y < 15*mm:
            p.showPage()
            y = height - 20*mm
            p.setFont('Helvetica', 8)
        row = [
            t.reference,
            t.from_warehouse.name[:25],
            t.to_warehouse.name[:25],
            str(t.items.count()),
            t.get_status_display(),
            t.created_at.strftime('%d/%m/%Y'),
            (t.created_by.display_name if t.created_by else '—')[:20],
        ]
        for x, val in zip(cols, row):
            p.drawString(x*mm, y, str(val))
        y -= 6*mm

    p.showPage()
    p.save()
    buffer.seek(0)
    response = HttpResponse(buffer, content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="mouvements_entrepots.pdf"'
    return response


@login_required
@seller_or_admin_required
def dash_transfers_export_excel(request):
    """Export Excel des transferts filtrés"""
    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font

    transfers = _filtered_transfers(request)[:1000]
    wb = Workbook()
    ws = wb.active
    ws.title = 'Mouvements'
    ws.append(['Référence', 'De', 'Vers', 'Produits', 'Quantité totale', 'Statut', 'Créé le', 'Par'])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for t in transfers:
        ws.append([
            t.reference,
            t.from_warehouse.name,
            t.to_warehouse.name,
            t.items.count(),
            sum(i.quantity for i in t.items.all()),
            t.get_status_display(),
            t.created_at.strftime('%d/%m/%Y %H:%M'),
            t.created_by.display_name if t.created_by else '',
        ])
    for col, w in zip('ABCDEFGH', [12, 25, 25, 10, 14, 12, 16, 18]):
        ws.column_dimensions[col].width = w

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="mouvements_entrepots.xlsx"'
    wb.save(response)
    return response


@login_required
@seller_or_admin_required
def dash_transfer_action(request, pk, action):
    """Confirmer / recevoir / annuler un transfert"""
    from inventory.models import StockTransfer
    transfer = get_object_or_404(StockTransfer, pk=pk)
    is_admin = request.user.is_superuser or request.user.role == 'admin'
    if not is_admin and transfer.store != request.user.store:
        django_messages.error(request, "Accès refusé.")
        return redirect('dashboard:transfers')

    if request.method == 'POST':
        if action == 'confirm' and transfer.confirm():
            django_messages.success(request, f'{transfer.reference} confirmé — stock sorti de {transfer.from_warehouse.name}.')
        elif action == 'receive' and transfer.receive():
            django_messages.success(request, f'{transfer.reference} reçu dans {transfer.to_warehouse.name}.')
        elif action == 'cancel' and transfer.cancel():
            django_messages.success(request, f'{transfer.reference} annulé.')
        else:
            django_messages.error(request, 'Action impossible dans cet état.')
    return redirect('dashboard:transfers')


@login_required
@seller_or_admin_required
def dash_stores(request):
    """Liste des boutiques de l'utilisateur"""
    from inventory.models import Warehouse
    stores = request.user.stores.all()
    warehouses = Warehouse.objects.filter(store__owner=request.user, is_active=True)
    return render(request, 'dashboard/stores.html', {
        'stores': stores,
        'warehouses': warehouses,
    })


# ======== EMPLOYÉS & RÔLES ========
@login_required
@seller_or_admin_required
def dash_employees(request):
    """Page d'accueil Employés : 3 cartes (Employés, Rôles, Paie)"""
    store = getattr(request.user, 'store', None)
    if not store:
        django_messages.error(request, "Vous devez avoir une boutique.")
        return redirect('dashboard:index')
    return render(request, 'dashboard/employees.html', {
        'members_count': store.members.filter(is_active=True).count(),
        'roles_count': store.roles.filter(is_active=True).count(),
        'payslips_count': store.members.filter(payslips__isnull=False).values('payslips').distinct().count(),
    })


@login_required
@seller_or_admin_required
def dash_employees_list(request):
    """Liste des employés avec ajout"""
    from store.models import StoreMember, StoreRole
    store = getattr(request.user, 'store', None)
    if not store:
        django_messages.error(request, "Vous devez avoir une boutique.")
        return redirect('dashboard:index')

    if request.method == 'POST':
        username = request.POST.get('username', '').strip()
        password = request.POST.get('password', '').strip()
        password_confirm = request.POST.get('password_confirm', '').strip()
        first_name = request.POST.get('first_name', '').strip()
        last_name = request.POST.get('last_name', '').strip()
        role_id = request.POST.get('role') or None
        warehouse_id = request.POST.get('warehouse') or None

        if not username:
            django_messages.error(request, "Le nom d'utilisateur est requis.")
            return redirect('dashboard:employees_list')

        user = User.objects.filter(username=username).first()
        if user is None:
            # Créer le compte employé avec mot de passe
            if not password or len(password) < 6:
                django_messages.error(request, "Mot de passe requis (min. 6 caractères) pour créer le compte.")
                return redirect('dashboard:employees_list')
            if password != password_confirm:
                django_messages.error(request, "Les mots de passe ne correspondent pas.")
                return redirect('dashboard:employees_list')
            user = User.objects.create_user(
                username=username, password=password,
                first_name=first_name, last_name=last_name,
                role='buyer',
            )
            # Envoyer les identifiants par email si l'employé a un email
            if user.email:
                from messaging.utils import notify
                notify(
                    user, 'account',
                    f'Vous avez rejoint {store.name}',
                    f'{request.user.display_name} vous a ajouté à son équipe.\n\nVos identifiants de connexion :\nNom d\'utilisateur : {username}\nMot de passe : {password}\n\nConnectez-vous pour accéder au tableau de bord.',
                    url='/compte/connexion/',
                    send_email=True,
                    email_subject=f'Vos identifiants {store.name} — AfriMarket',
                )
            django_messages.success(request, f'Compte créé pour {user.display_name}.')
        elif user == request.user:
            django_messages.error(request, "Vous êtes déjà le propriétaire.")
            return redirect('dashboard:employees_list')

        salary = int(request.POST.get('salary', 0) or 0)
        member, created = StoreMember.objects.get_or_create(
            store=store, user=user,
            defaults={'role_id': role_id, 'warehouse_id': warehouse_id, 'salary': salary}
        )
        if not created:
            member.role_id = role_id
            member.warehouse_id = warehouse_id
            member.salary = salary
            member.is_active = True
            member.save()
        django_messages.success(request, f'{user.display_name} fait maintenant partie de votre équipe.')
        return redirect('dashboard:employees_list')

    members = store.members.select_related('user', 'role', 'warehouse')
    warehouse_id = request.GET.get('warehouse', '')
    if warehouse_id:
        members = members.filter(warehouse_id=warehouse_id)
    search_query = request.GET.get('q', '').strip()
    if search_query:
        from django.db.models import Q
        members = members.filter(
            Q(user__username__icontains=search_query) |
            Q(user__first_name__icontains=search_query) |
            Q(user__last_name__icontains=search_query)
        )

    return render(request, 'dashboard/employees_list.html', {
        'members': members,
        'owner': store.owner,
        'roles': store.roles.filter(is_active=True),
        'warehouses': store.warehouses.filter(is_active=True),
        'warehouse_filter': warehouse_id,
        'search_query': search_query,
    })


@login_required
@seller_or_admin_required
def dash_employee_toggle(request, pk):
    """Activer/désactiver un employé"""
    from store.models import StoreMember
    member = get_object_or_404(StoreMember, pk=pk, store=request.user.store)
    if request.method == 'POST':
        member.is_active = not member.is_active
        member.save()
        django_messages.success(request, f'{member.user.display_name} {"activé" if member.is_active else "désactivé"}.')
    return redirect('dashboard:employees_list')


@login_required
@seller_or_admin_required
def dash_roles(request):
    """Gestion des rôles et permissions"""
    from store.models import StoreRole
    store = getattr(request.user, 'store', None)
    if not store:
        django_messages.error(request, "Vous devez avoir une boutique.")
        return redirect('dashboard:index')

    if request.method == 'POST':
        role_id = request.POST.get('role_id')
        permissions = request.POST.getlist('permissions')
        if role_id:
            role = get_object_or_404(StoreRole, pk=role_id, store=store)
            role.name = request.POST.get('name', role.name)
            role.permissions = permissions
            role.save()
            django_messages.success(request, f'Rôle "{role.name}" mis à jour.')
        else:
            StoreRole.objects.create(
                store=store,
                name=request.POST.get('name'),
                permissions=permissions,
            )
            django_messages.success(request, 'Rôle créé !')
        return redirect('dashboard:roles')

    return render(request, 'dashboard/roles.html', {
        'roles': store.roles.all(),
        'permission_choices': StoreRole.PERMISSION_CHOICES,
    })


@login_required
@seller_or_admin_required
def dash_employee_detail(request, pk):
    """Détail d'un employé : infos + bulletins de paie"""
    from store.models import StoreMember
    member = get_object_or_404(StoreMember, pk=pk, store=request.user.store)
    return render(request, 'dashboard/employee_detail.html', {
        'member': member,
        'payslips': member.payslips.all(),
    })


# ======== PAIE DES EMPLOYÉS ========
@login_required
@seller_or_admin_required
def dash_payroll(request):
    """Liste des bulletins de paie"""
    from store.models import Payslip
    store = getattr(request.user, 'store', None)
    if not store:
        django_messages.error(request, "Vous devez avoir une boutique.")
        return redirect('dashboard:index')

    payslips = Payslip.objects.filter(member__store=store).select_related('member', 'member__user')

    month = request.GET.get('month', '')
    year = request.GET.get('year', '')
    if month:
        payslips = payslips.filter(month=month)
    if year:
        payslips = payslips.filter(year=year)

    return render(request, 'dashboard/payroll.html', {
        'payslips': payslips,
        'month_filter': month,
        'year_filter': year,
        'months': [(i, timezone.datetime(2000, i, 1).strftime('%B')) for i in range(1, 13)],
        'years': range(timezone.now().year - 2, timezone.now().year + 1),
        'members': store.members.filter(is_active=True).select_related('user'),
    })


@login_required
@seller_or_admin_required
def dash_payslip_create(request):
    """Créer un bulletin de paie"""
    from store.models import Payslip, StoreMember
    store = getattr(request.user, 'store', None)
    if not store:
        return redirect('dashboard:index')

    if request.method == 'POST':
        from store.models import PayslipLine
        member = get_object_or_404(StoreMember, pk=request.POST.get('member'), store=store)
        month = int(request.POST.get('month'))
        year = int(request.POST.get('year'))
        # Vérifier si déjà existant
        if Payslip.objects.filter(member=member, month=month, year=year).exists():
            django_messages.error(request, 'Un bulletin existe déjà pour cet employé ce mois-ci.')
            return redirect('dashboard:payroll')
        payslip = Payslip.objects.create(
            member=member,
            month=month,
            year=year,
            base_salary=int(request.POST.get('base_salary', member.salary or 0)),
            notes=request.POST.get('notes', ''),
            created_by=request.user,
        )
        # Lignes de primes et retenues
        line_types = request.POST.getlist('line_type[]')
        line_labels = request.POST.getlist('line_label[]')
        line_amounts = request.POST.getlist('line_amount[]')
        for i, ltype in enumerate(line_types):
            label = line_labels[i].strip() if i < len(line_labels) else ''
            amount = int(line_amounts[i]) if i < len(line_amounts) and line_amounts[i] else 0
            if label and amount > 0:
                PayslipLine.objects.create(payslip=payslip, line_type=ltype, label=label, amount=amount)
        django_messages.success(request, f'Bulletin {payslip.reference} créé !')
    return redirect('dashboard:payroll')


@login_required
@seller_or_admin_required
def dash_payslip_action(request, pk, action):
    """Valider / marquer payé un bulletin"""
    from store.models import Payslip
    payslip = get_object_or_404(Payslip, pk=pk, member__store=request.user.store)
    if request.method == 'POST':
        if action == 'validate' and payslip.status == 'draft':
            payslip.status = 'validated'
            payslip.save()
            django_messages.success(request, f'{payslip.reference} validé.')
        elif action == 'pay' and payslip.status == 'validated':
            payslip.mark_paid()
            django_messages.success(request, f'{payslip.reference} marqué comme payé.')
    return redirect('dashboard:payroll')


# ======== GESTION DE STOCK ========
@login_required
@seller_or_admin_required
def dash_inventory(request):
    """Page inventaire : historique des mouvements de stock avec filtres"""
    from catalog.models import StockMovement
    from datetime import datetime
    if request.user.is_seller and hasattr(request.user, 'store'):
        movements = StockMovement.objects.filter(store=request.user.store)
    else:
        movements = StockMovement.objects.all()

    q = request.GET.get('q', '').strip()
    if q:
        movements = movements.filter(product__name__icontains=q)
    warehouse_id = request.GET.get('warehouse')
    current_warehouse = None
    if warehouse_id:
        from inventory.models import Warehouse
        movements = movements.filter(warehouse_id=warehouse_id)
        current_warehouse = Warehouse.objects.filter(pk=warehouse_id).first()
    mtype = request.GET.get('type', '')
    if mtype:
        movements = movements.filter(movement_type=mtype)
    date_from = request.GET.get('from', '')
    date_to = request.GET.get('to', '')
    if date_from:
        try:
            movements = movements.filter(created_at__date__gte=datetime.strptime(date_from, '%Y-%m-%d').date())
        except ValueError:
            pass
    if date_to:
        try:
            movements = movements.filter(created_at__date__lte=datetime.strptime(date_to, '%Y-%m-%d').date())
        except ValueError:
            pass

    # Résumé par produit : ventes, entrées, sorties, stock actuel
    from catalog.models import Product
    product_stats = []
    if current_warehouse:
        product_ids = movements.values_list('product_id', flat=True).distinct()
        for pid in product_ids:
            product = Product.objects.filter(pk=pid).first()
            if not product:
                continue
            mvts = movements.filter(product_id=pid)
            sales = mvts.filter(movement_type='sale').aggregate(t=Sum('quantity'))['t'] or 0
            entries = mvts.filter(movement_type='in').aggregate(t=Sum('quantity'))['t'] or 0
            exits = mvts.filter(movement_type='out').aggregate(t=Sum('quantity'))['t'] or 0
            stock = product.warehouse_stocks.filter(warehouse=current_warehouse).first()
            product_stats.append({
                'product': product,
                'sales': abs(sales),
                'entries': entries,
                'exits': abs(exits),
                'current_stock': stock.quantity if stock else 0,
            })

    return render(request, 'dashboard/inventory.html', {
        'movements': movements.select_related('product', 'created_by')[:100],
        'search_query': q,
        'warehouse': current_warehouse,
        'type_filter': mtype,
        'date_from': date_from,
        'date_to': date_to,
        'type_choices': StockMovement.TYPE_CHOICES,
        'total_movements': movements.count(),
        'product_stats': product_stats,
    })


@login_required
@seller_or_admin_required
def dash_stock_adjust(request, pk):
    """Ajustement / entrée / sortie de stock via modale"""
    product = get_object_or_404(Product, pk=pk)
    is_admin = request.user.is_superuser or request.user.role == 'admin'
    if not is_admin and product.store != request.user.store:
        django_messages.error(request, "Ce produit ne vous appartient pas.")
        return redirect('dashboard:products')

    if request.method == 'POST':
        movement_type = request.POST.get('movement_type', 'adjustment')
        qty = int(request.POST.get('quantity', 0))
        reason = request.POST.get('reason', '').strip()
        threshold = request.POST.get('low_stock_threshold')

        if threshold is not None and threshold != '':
            product.low_stock_threshold = max(0, int(threshold))
            product.save(update_fields=['low_stock_threshold'])

        if qty > 0:
            if movement_type == 'out':
                qty = -qty
            product.adjust_stock(qty, movement_type, user=request.user, reason=reason)
            django_messages.success(request, f'Stock de "{product.name}" mis à jour : {product.stock} unités.')
        else:
            django_messages.error(request, 'La quantité doit être supérieure à 0.')

    return redirect(request.META.get('HTTP_REFERER', 'dashboard:products'))


def _filtered_movements(request):
    """Retourne les mouvements filtrés selon les paramètres GET."""
    from catalog.models import StockMovement
    from datetime import datetime
    if request.user.is_seller and hasattr(request.user, 'store'):
        movements = StockMovement.objects.filter(store=request.user.store)
    else:
        movements = StockMovement.objects.all()
    q = request.GET.get('q', '').strip()
    if q:
        movements = movements.filter(product__name__icontains=q)
    mtype = request.GET.get('type', '')
    if mtype:
        movements = movements.filter(movement_type=mtype)
    date_from = request.GET.get('from', '')
    date_to = request.GET.get('to', '')
    if date_from:
        try:
            movements = movements.filter(created_at__date__gte=datetime.strptime(date_from, '%Y-%m-%d').date())
        except ValueError:
            pass
    if date_to:
        try:
            movements = movements.filter(created_at__date__lte=datetime.strptime(date_to, '%Y-%m-%d').date())
        except ValueError:
            pass
    return movements.select_related('product', 'created_by')


@login_required
@seller_or_admin_required
def dash_movements_export_pdf(request):
    """Export PDF de la liste filtrée des mouvements"""
    from django.http import HttpResponse
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas
    import io

    movements = _filtered_movements(request)[:500]

    buffer = io.BytesIO()
    p = canvas.Canvas(buffer, pagesize=landscape(A4))
    width, height = landscape(A4)

    p.setFont('Helvetica-Bold', 16)
    p.drawString(15*mm, height - 18*mm, 'AfriMarket — Mouvements de stock')
    p.setFont('Helvetica', 9)
    p.drawString(15*mm, height - 24*mm, f'Généré le {timezone.now().strftime("%d/%m/%Y à %H:%M")} — {movements.count()} mouvement(s)')
    p.line(15*mm, height - 27*mm, width - 15*mm, height - 27*mm)

    y = height - 36*mm
    p.setFont('Helvetica-Bold', 8)
    headers = ['Date', 'Produit', 'Type', 'Qté', 'Avant', 'Après', 'Motif', 'Référence', 'Par']
    cols = [15, 40, 95, 125, 140, 155, 170, 215, 245]
    for x, h in zip(cols, headers):
        p.drawString(x*mm, y, h)
    p.line(15*mm, y - 2*mm, width - 15*mm, y - 2*mm)
    y -= 8*mm

    p.setFont('Helvetica', 8)
    for m in movements:
        if y < 15*mm:
            p.showPage()
            y = height - 20*mm
            p.setFont('Helvetica', 8)
        row = [
            m.created_at.strftime('%d/%m/%y %H:%M'),
            m.product.name[:30],
            m.get_movement_type_display(),
            f'{m.quantity:+d}',
            str(m.stock_before),
            str(m.stock_after),
            (m.reason or '—')[:25],
            (m.reference or '—')[:15],
            (m.created_by.display_name if m.created_by else '—')[:15],
        ]
        for x, val in zip(cols, row):
            p.drawString(x*mm, y, str(val))
        y -= 6*mm

    p.showPage()
    p.save()
    buffer.seek(0)
    response = HttpResponse(buffer, content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="mouvements_stock.pdf"'
    return response


@login_required
@seller_or_admin_required
def dash_movements_export_excel(request):
    """Export Excel de la liste filtrée des mouvements"""
    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font

    movements = _filtered_movements(request)[:1000]

    wb = Workbook()
    ws = wb.active
    ws.title = 'Mouvements'
    headers = ['Date', 'Produit', 'Type', 'Quantité', 'Stock avant', 'Stock après', 'Motif', 'Référence', 'Effectué par']
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for m in movements:
        ws.append([
            m.created_at.strftime('%d/%m/%Y %H:%M'),
            m.product.name,
            m.get_movement_type_display(),
            m.quantity,
            m.stock_before,
            m.stock_after,
            m.reason or '',
            m.reference or '',
            m.created_by.display_name if m.created_by else '',
        ])
    for col, w in zip('ABCDEFGHI', [16, 30, 14, 10, 12, 12, 30, 15, 15]):
        ws.column_dimensions[col].width = w

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="mouvements_stock.xlsx"'
    wb.save(response)
    return response


@login_required
@seller_or_admin_required
def dash_movement_pdf(request, pk):
    """Export PDF d'un mouvement de stock"""
    from catalog.models import StockMovement
    from django.http import HttpResponse
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas
    import io

    m = get_object_or_404(StockMovement, pk=pk)
    is_admin = request.user.is_superuser or request.user.role == 'admin'
    if not is_admin and m.store != request.user.store:
        django_messages.error(request, "Accès refusé.")
        return redirect('dashboard:inventory')

    buffer = io.BytesIO()
    p = canvas.Canvas(buffer, pagesize=A4)
    width, height = A4

    p.setFont('Helvetica-Bold', 20)
    p.drawString(20*mm, height - 25*mm, 'AfriMarket — Mouvement de stock')
    p.setLineWidth(1)
    p.line(20*mm, height - 30*mm, width - 20*mm, height - 30*mm)

    rows = [
        ('Référence mouvement', f'#{m.id}'),
        ('Produit', m.product.name),
        ('Type', m.get_movement_type_display()),
        ('Quantité', f'{m.quantity:+d}'),
        ('Stock avant', str(m.stock_before)),
        ('Stock après', str(m.stock_after)),
        ('Motif', m.reason or '—'),
        ('Référence', m.reference or '—'),
        ('Effectué par', m.created_by.display_name if m.created_by else '—'),
        ('Date', m.created_at.strftime('%d/%m/%Y %H:%M')),
    ]
    y = height - 45*mm
    for label, value in rows:
        p.setFont('Helvetica', 10)
        p.drawString(25*mm, y, label)
        p.setFont('Helvetica-Bold', 11)
        p.drawString(80*mm, y, str(value))
        y -= 10*mm

    p.setFont('Helvetica', 8)
    p.drawString(20*mm, 15*mm, f'Document généré le {timezone.now().strftime("%d/%m/%Y à %H:%M")} — AfriMarket')
    p.showPage()
    p.save()
    buffer.seek(0)

    response = HttpResponse(buffer, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="mouvement_{m.id}.pdf"'
    return response


@login_required
@seller_or_admin_required
def dash_movement_excel(request, pk):
    """Export Excel d'un mouvement de stock"""
    from catalog.models import StockMovement
    from django.http import HttpResponse
    from openpyxl import Workbook

    m = get_object_or_404(StockMovement, pk=pk)
    is_admin = request.user.is_superuser or request.user.role == 'admin'
    if not is_admin and m.store != request.user.store:
        django_messages.error(request, "Accès refusé.")
        return redirect('dashboard:inventory')

    wb = Workbook()
    ws = wb.active
    ws.title = 'Mouvement'
    ws.append(['Champ', 'Valeur'])
    for cell in ws[1]:
        cell.font = cell.font.copy(bold=True)
    rows = [
        ('Référence mouvement', f'#{m.id}'),
        ('Produit', m.product.name),
        ('Type', m.get_movement_type_display()),
        ('Quantité', m.quantity),
        ('Stock avant', m.stock_before),
        ('Stock après', m.stock_after),
        ('Motif', m.reason or '—'),
        ('Référence', m.reference or '—'),
        ('Effectué par', m.created_by.display_name if m.created_by else '—'),
        ('Date', m.created_at.strftime('%d/%m/%Y %H:%M')),
    ]
    for row in rows:
        ws.append(list(row))
    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 40

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="mouvement_{m.id}.xlsx"'
    wb.save(response)
    return response


@login_required
@seller_or_admin_required
def dash_customers(request):
    # Si l'utilisateur est un vendeur, afficher uniquement ses clients
    if request.user.is_seller and hasattr(request.user, 'store'):
        customers = User.objects.filter(
            orders__items__product__store__owner=request.user
        ).distinct().annotate(
            order_count=Count('orders', filter=Q(orders__items__product__store__owner=request.user), distinct=True),
            total_spent=Sum('orders__total_amount', filter=Q(orders__items__product__store__owner=request.user))
        ).order_by('-date_joined')
    else:
        customers = User.objects.filter(role='buyer').annotate(
            order_count=Count('orders'), total_spent=Sum('orders__total_amount')
        ).order_by('-date_joined')

    # Recherche
    q = request.GET.get('q', '').strip()
    if q:
        customers = customers.filter(
            Q(username__icontains=q) | Q(first_name__icontains=q) |
            Q(last_name__icontains=q) | Q(email__icontains=q) | Q(phone__icontains=q)
        )

    customers_list = list(customers)
    total_spent_all = sum(c.total_spent or 0 for c in customers_list)
    total_orders_all = sum(c.order_count or 0 for c in customers_list)

    return render(request, 'dashboard/customers.html', {
        'customers': customers_list,
        'total': len(customers_list),
        'new_30d': sum(1 for c in customers_list if c.date_joined >= timezone.now() - timedelta(days=30)),
        'total_spent_all': total_spent_all,
        'avg_spent': int(total_spent_all / len(customers_list)) if customers_list else 0,
        'search_query': q,
    })


@login_required
@seller_or_admin_required
def dash_customers_export(request):
    """Export Excel de la liste des clients"""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from django.http import HttpResponse

    if request.user.is_seller and hasattr(request.user, 'store'):
        customers = User.objects.filter(
            orders__items__product__store__owner=request.user
        ).distinct().annotate(
            order_count=Count('orders', filter=Q(orders__items__product__store__owner=request.user), distinct=True),
            total_spent=Sum('orders__total_amount', filter=Q(orders__items__product__store__owner=request.user))
        )
    else:
        customers = User.objects.filter(role='buyer').annotate(
            order_count=Count('orders'), total_spent=Sum('orders__total_amount')
        )

    q = request.GET.get('q', '').strip()
    if q:
        customers = customers.filter(
            Q(username__icontains=q) | Q(first_name__icontains=q) |
            Q(last_name__icontains=q) | Q(email__icontains=q) | Q(phone__icontains=q)
        )

    wb = Workbook()
    ws = wb.active
    ws.title = 'Clients'
    ws.append(['Nom', 'Email', 'Téléphone', 'Ville', 'Commandes', 'Total dépensé (F)', 'Inscrit le'])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for c in customers:
        ws.append([
            c.display_name, c.email, c.phone or '', c.city or '',
            c.order_count or 0, int(c.total_spent or 0),
            c.date_joined.strftime('%d/%m/%Y'),
        ])
    for col, w in zip('ABCDEFG', [25, 30, 15, 15, 12, 18, 12]):
        ws.column_dimensions[col].width = w

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="clients.xlsx"'
    wb.save(response)
    return response


@login_required
@seller_or_admin_required
def dash_store(request):
    store = getattr(request.user, 'store', None)
    if not store:
        django_messages.error(request, 'Vous n\'avez pas de boutique.')
        return redirect('dashboard:index')

    # Créer une nouvelle boutique
    if request.method == 'POST' and request.POST.get('action') == 'create_store':
        from inventory.models import Warehouse
        new_store = Store.objects.create(
            owner=request.user,
            name=request.POST.get('name'),
            city=request.POST.get('city', 'Douala'),
            phone=request.POST.get('phone', ''),
            is_active='is_active' in request.POST,
        )
        # Lier à l'entrepôt choisi
        warehouse_id = request.POST.get('warehouse')
        if warehouse_id:
            wh = Warehouse.objects.filter(pk=warehouse_id, store__owner=request.user).first()
            if wh:
                wh.linked_stores.add(new_store)
        django_messages.success(request, f'Boutique "{new_store.name}" créée !')
        return redirect('dashboard:stores')

    if request.method == 'POST':
        store.name = request.POST.get('name', store.name)
        store.description = request.POST.get('description', '')
        store.phone = request.POST.get('phone', '')
        store.whatsapp = request.POST.get('whatsapp', '')
        store.email = request.POST.get('email', '')
        store.address = request.POST.get('address', '')
        store.city = request.POST.get('city', '')
        store.latitude = request.POST.get('latitude') or None
        store.longitude = request.POST.get('longitude') or None
        store.opening_hours = request.POST.get('opening_hours', '')
        store.facebook = request.POST.get('facebook', '')
        store.instagram = request.POST.get('instagram', '')
        if request.FILES.get('logo'): store.logo = request.FILES['logo']
        if request.FILES.get('banner'): store.banner = request.FILES['banner']
        store.save()
        django_messages.success(request, 'Boutique mise à jour !')

    # Score de complétion du profil
    fields = [store.name, store.description, store.phone, store.email, store.address,
              store.city, store.logo, store.banner, store.whatsapp, store.opening_hours]
    completion = int(sum(1 for f in fields if f) / len(fields) * 100)

    return render(request, 'dashboard/store.html', {
        'store': store,
        'completion': completion,
        'products_count': store.products.count(),
        'total_sales': store.total_sales,
        'rating': store.rating,
    })


@login_required
@seller_or_admin_required
def dash_notifications(request):
    """Centre de notifications du dashboard"""
    from messaging.models import Notification
    notifs = Notification.objects.filter(user=request.user)
    notifs.filter(is_read=False).update(is_read=True)
    return render(request, 'dashboard/notifications.html', {'notifications': notifs[:100]})


@login_required
@seller_or_admin_required
def dash_settings(request):
    """Page de paramétrage : boutique, facturation, compte"""
    from invoicing.models import InvoiceSettings
    store = getattr(request.user, 'store', None)
    if not store:
        django_messages.error(request, "Vous devez avoir une boutique.")
        return redirect('dashboard:index')

    invoice_settings, _ = InvoiceSettings.objects.get_or_create(
        store=store,
        defaults={
            'company_name': store.name,
            'address': store.address,
            'city': store.city,
            'phone': store.phone,
            'email': store.email,
        }
    )

    if request.method == 'POST':
        section = request.POST.get('section')

        if section == 'invoicing':
            s = invoice_settings
            s.company_name = request.POST.get('company_name', '')
            s.tax_id = request.POST.get('tax_id', '')
            s.registration_number = request.POST.get('registration_number', '')
            s.address = request.POST.get('address', '')
            s.city = request.POST.get('city', '')
            s.country = request.POST.get('country', 'Cameroun')
            s.phone = request.POST.get('phone', '')
            s.email = request.POST.get('email', '')
            s.website = request.POST.get('website', '')
            s.apply_tva = 'apply_tva' in request.POST
            s.tva_rate = request.POST.get('tva_rate', 19.25)
            s.invoice_prefix = request.POST.get('invoice_prefix', 'FAC')
            s.header_text = request.POST.get('header_text', '')
            s.footer_text = request.POST.get('footer_text', '')
            s.payment_terms = request.POST.get('payment_terms', '')
            s.bank_details = request.POST.get('bank_details', '')
            s.signature_name = request.POST.get('signature_name', '')
            s.signature_title = request.POST.get('signature_title', '')
            if request.FILES.get('logo'):
                s.logo = request.FILES['logo']
            if request.FILES.get('signature'):
                s.signature = request.FILES['signature']
            s.save()
            django_messages.success(request, 'Paramètres de facturation enregistrés.')

        elif section == 'notifications':
            from messaging.models import NotificationPreference
            prefs, _ = NotificationPreference.objects.get_or_create(user=request.user)
            prefs.internal_enabled = 'internal_enabled' in request.POST
            prefs.email_order = 'email_order' in request.POST
            prefs.email_rfq = 'email_rfq' in request.POST
            prefs.email_invoice = 'email_invoice' in request.POST
            prefs.email_payment = 'email_payment' in request.POST
            prefs.email_stock = 'email_stock' in request.POST
            prefs.email_account = 'email_account' in request.POST
            prefs.save()
            django_messages.success(request, 'Préférences de notification enregistrées.')

        elif section == 'account':
            user = request.user
            user.first_name = request.POST.get('first_name', '')
            user.last_name = request.POST.get('last_name', '')
            user.email = request.POST.get('email', '')
            user.phone = request.POST.get('phone', '')
            user.save()
            new_password = request.POST.get('new_password', '').strip()
            if new_password:
                if not user.check_password(request.POST.get('current_password', '')):
                    django_messages.error(request, 'Mot de passe actuel incorrect.')
                    return redirect('dashboard:settings')
                if len(new_password) < 6:
                    django_messages.error(request, 'Le nouveau mot de passe doit faire au moins 6 caractères.')
                    return redirect('dashboard:settings')
                if new_password != request.POST.get('confirm_password', ''):
                    django_messages.error(request, 'Les mots de passe ne correspondent pas.')
                    return redirect('dashboard:settings')
                user.set_password(new_password)
                user.save()
                from django.contrib.auth import update_session_auth_hash
                update_session_auth_hash(request, user)
                django_messages.success(request, 'Compte et mot de passe mis à jour.')
            else:
                django_messages.success(request, 'Informations du compte enregistrées.')

        return redirect('dashboard:settings')

    from messaging.models import NotificationPreference
    notif_prefs, _ = NotificationPreference.objects.get_or_create(user=request.user)

    return render(request, 'dashboard/settings.html', {
        'store': store,
        'invoice_settings': invoice_settings,
        'notif_prefs': notif_prefs,
    })


# ======== ADMIN USER MANAGEMENT ========
@login_required
def admin_users(request):
    if not request.user.is_superuser and request.user.role != 'admin':
        return redirect('dashboard:index')
    users = User.objects.all().annotate(
        order_count=Count('orders', distinct=True)
    ).order_by('-date_joined')
    role = request.GET.get('role')
    if role: users = users.filter(role=role)
    search = request.GET.get('q', '')
    if search: users = users.filter(
        Q(username__icontains=search) | Q(email__icontains=search) |
        Q(first_name__icontains=search) | Q(last_name__icontains=search)
    )
    return render(request, 'dashboard/users.html', {
        'users': users,
        'total': User.objects.count(),
        'buyers': User.objects.filter(role='buyer').count(),
        'sellers': User.objects.filter(role='seller').count(),
        'admins': User.objects.filter(role='admin').count(),
        'verified': User.objects.filter(is_verified=True).count(),
        'search': search, 'current_role': role,
    })


@login_required
def admin_user_detail(request, pk):
    if not request.user.is_superuser: return redirect('dashboard:index')
    target = get_object_or_404(User, pk=pk)
    orders = Order.objects.filter(buyer=target) if target.role == 'buyer' else Order.objects.none()
    store = getattr(target, 'store', None) if target.is_seller else None
    return render(request, 'dashboard/user_detail.html', {
        'target': target, 'orders': orders, 'store': store,
    })


@login_required
def admin_user_action(request, pk):
    if not request.user.is_superuser: return redirect('dashboard:index')
    target = get_object_or_404(User, pk=pk)
    if request.method == 'POST':
        action = request.POST.get('action')
        if action == 'verify':
            target.is_verified = True
            target.save()
            django_messages.success(request, f'{target.display_name} vérifié.')
        elif action == 'unverify':
            target.is_verified = False
            target.save()
        elif action == 'deactivate':
            target.is_active = False
            target.save()
            django_messages.warning(request, f'{target.display_name} désactivé.')
        elif action == 'activate':
            target.is_active = True
            target.save()
            django_messages.success(request, f'{target.display_name} réactivé.')
        elif action == 'make_admin':
            target.role = 'admin'
            target.is_staff = True
            target.save()
            django_messages.success(request, f'{target.display_name} est maintenant admin.')
        elif action == 'change_role':
            new_role = request.POST.get('new_role')
            if new_role: target.role = new_role; target.save()
    return redirect('dashboard:user_detail', pk=pk)


@login_required
def admin_stores(request):
    if not request.user.is_superuser: return redirect('dashboard:index')
    stores = Store.objects.annotate(
        product_cnt=Count('products'),
        revenue=Sum(F('products__orderitem__price') * F('products__orderitem__quantity'),
                     filter=Q(products__orderitem__order__is_paid=True))
    ).order_by('-created_at')
    return render(request, 'dashboard/stores.html', {'stores': stores})


@login_required
def admin_verify_store(request, pk):
    if not request.user.is_superuser: return redirect('dashboard:index')
    store = get_object_or_404(Store, pk=pk)
    store.is_verified = not store.is_verified
    store.save()
    status = 'vérifiée' if store.is_verified else 'non vérifiée'
    django_messages.success(request, f'Boutique {store.name} : {status}.')
    return redirect('dashboard:stores')


@login_required
def analytics(request):
    if not request.user.is_superuser: return redirect('dashboard:index')
    now = timezone.now()
    # Conversion funnel
    total_visits = Product.objects.aggregate(t=Sum('views_count'))['t'] or 1
    total_cart = Order.objects.count()
    total_paid = Order.objects.filter(is_paid=True).count()
    # Growth
    this_month = Order.objects.filter(created_at__month=now.month, created_at__year=now.year, is_paid=True).aggregate(t=Sum('total_amount'))['t'] or 0
    last_month = Order.objects.filter(created_at__month=(now.month - 1) or 12, is_paid=True).aggregate(t=Sum('total_amount'))['t'] or 1
    growth = round((this_month - last_month) / max(last_month, 1) * 100, 1)
    # Category performance
    cat_perf = Category.objects.filter(parent__isnull=True).annotate(
        rev=Sum(F('products__orderitem__price') * F('products__orderitem__quantity'),
                filter=Q(products__orderitem__order__is_paid=True)),
        cnt=Count('products__orderitem', filter=Q(products__orderitem__order__is_paid=True))
    ).order_by('-rev')[:8]
    cat_labels = [c.name for c in cat_perf]
    cat_data = [int(c.rev or 0) for c in cat_perf]
    return render(request, 'dashboard/analytics.html', {
        'total_visits': total_visits, 'total_cart': total_cart,
        'total_paid': total_paid, 'growth': growth,
        'this_month': this_month, 'last_month': last_month,
        'cat_labels': json.dumps(cat_labels), 'cat_data': json.dumps(cat_data),
    })


@login_required
def admin_users(request):
    if not (request.user.is_superuser or request.user.role == 'admin'):
        return redirect('dashboard:index')
    users = User.objects.all().order_by('-date_joined')
    role = request.GET.get('role')
    if role:
        users = users.filter(role=role)
    search = request.GET.get('q', '')
    if search:
        users = users.filter(
            Q(username__icontains=search) | Q(email__icontains=search) |
            Q(first_name__icontains=search) | Q(last_name__icontains=search)
        )
    return render(request, 'dashboard/users.html', {
        'users': users, 'total': users.count(),
        'buyers': User.objects.filter(role='buyer').count(),
        'sellers_count': User.objects.filter(role='seller').count(),
        'admins': User.objects.filter(role='admin').count(),
    })


@login_required
def admin_user_detail(request, pk):
    if not (request.user.is_superuser or request.user.role == 'admin'):
        return redirect('dashboard:index')
    u = get_object_or_404(User, pk=pk)
    orders_qs = Order.objects.filter(buyer=u) if u.role == 'buyer' else Order.objects.filter(items__store__owner=u).distinct()
    from accounting.models import Transaction, SellerWallet
    txns = Transaction.objects.filter(user=u)[:20]
    wallet = SellerWallet.objects.filter(user=u).first() if u.is_seller else None
    from billing.models import Subscription
    sub = Subscription.objects.filter(user=u).first()
    return render(request, 'dashboard/user_detail.html', {
        'profile_user': u, 'orders': orders_qs[:10],
        'transactions': txns, 'wallet': wallet, 'subscription': sub,
    })


@login_required
def admin_toggle_user(request, pk):
    if not (request.user.is_superuser or request.user.role == 'admin'):
        return redirect('dashboard:index')
    u = get_object_or_404(User, pk=pk)
    u.is_active = not u.is_active
    u.save()
    status = 'activé' if u.is_active else 'désactivé'
    django_messages.success(request, f'Utilisateur {u.username} {status}.')
    return redirect('dashboard:users')


@login_required
def admin_sellers(request):
    if not (request.user.is_superuser or request.user.role == 'admin'):
        return redirect('dashboard:index')
    sellers = Store.objects.select_related('owner').all()
    return render(request, 'dashboard/sellers.html', {'sellers': sellers})


@login_required
def admin_verify_seller(request, pk):
    if not (request.user.is_superuser or request.user.role == 'admin'):
        return redirect('dashboard:index')
    store = get_object_or_404(Store, pk=pk)
    store.is_verified = not store.is_verified
    store.save()
    store.owner.is_verified = store.is_verified
    store.owner.save()
    status = 'vérifié' if store.is_verified else 'non vérifié'
    django_messages.success(request, f'{store.name} est maintenant {status}.')
    return redirect('dashboard:sellers')


@login_required
def admin_payouts(request):
    if not (request.user.is_superuser or request.user.role == 'admin'):
        return redirect('dashboard:index')
    from accounting.models import PayoutRequest
    payouts = PayoutRequest.objects.select_related('user').all()
    status = request.GET.get('status')
    if status:
        payouts = payouts.filter(status=status)
    return render(request, 'dashboard/payouts.html', {'payouts': payouts})


@login_required
def admin_process_payout(request, pk):
    if not (request.user.is_superuser or request.user.role == 'admin'):
        return redirect('dashboard:index')
    from accounting.models import PayoutRequest, SellerWallet, Transaction
    payout = get_object_or_404(PayoutRequest, pk=pk)
    if request.method == 'POST':
        action = request.POST.get('action')
        if action == 'approve':
            payout.status = 'completed'
            payout.processed_at = timezone.now()
            payout.save()
            wallet = SellerWallet.objects.get(user=payout.user)
            wallet.pending_balance -= payout.amount
            wallet.total_withdrawn += payout.amount
            wallet.save()
            Transaction.objects.create(
                user=payout.user, type='payout', amount=payout.amount,
                status='completed', reference=f'Versement #{payout.pk}',
            )
            django_messages.success(request, f'Versement de {payout.amount:,.0f} F approuvé.')
        elif action == 'reject':
            payout.status = 'rejected'
            payout.admin_notes = request.POST.get('notes', '')
            payout.save()
            wallet = SellerWallet.objects.get(user=payout.user)
            wallet.pending_balance -= payout.amount
            wallet.balance += payout.amount
            wallet.save()
            django_messages.info(request, 'Versement rejeté.')
    return redirect('dashboard:payouts')


@login_required
def admin_platform_stats(request):
    if not (request.user.is_superuser or request.user.role == 'admin'):
        return redirect('dashboard:index')
    from accounting.models import Transaction, Expense
    from billing.models import Subscription, ProductBoost

    total_users = User.objects.count()
    active_subscriptions = Subscription.objects.filter(status='active').count()
    total_sub_revenue = Transaction.objects.filter(type='subscription', status='completed').aggregate(t=Sum('amount'))['t'] or 0
    total_boost_revenue = Transaction.objects.filter(type='boost', status='completed').aggregate(t=Sum('amount'))['t'] or 0
    total_commissions = Transaction.objects.filter(type='commission', status='completed').aggregate(t=Sum('amount'))['t'] or 0
    total_expenses = Expense.objects.aggregate(t=Sum('amount'))['t'] or 0
    total_orders = Order.objects.count()
    gmv = Order.objects.filter(is_paid=True).aggregate(t=Sum('total_amount'))['t'] or 0

    return render(request, 'dashboard/platform_stats.html', {
        'total_users': total_users,
        'active_subscriptions': active_subscriptions,
        'total_sub_revenue': total_sub_revenue,
        'total_boost_revenue': total_boost_revenue,
        'total_commissions': total_commissions,
        'total_expenses': total_expenses,
        'total_orders': total_orders,
        'gmv': gmv,
        'net_revenue': (total_sub_revenue or 0) + (total_boost_revenue or 0) + (total_commissions or 0) - (total_expenses or 0),
    })


@login_required
def admin_banners(request):
    if not (request.user.is_superuser or request.user.role == 'admin'):
        return redirect('dashboard:index')
    from catalog.models import HeroBanner
    banners = HeroBanner.objects.all()
    return render(request, 'dashboard/banners.html', {'banners': banners})


@login_required
def admin_banner_edit(request, pk=None):
    if not (request.user.is_superuser or request.user.role == 'admin'):
        return redirect('dashboard:index')
    from catalog.models import HeroBanner
    banner = get_object_or_404(HeroBanner, pk=pk) if pk else None
    if request.method == 'POST':
        data = {
            'title': request.POST.get('title', ''),
            'subtitle': request.POST.get('subtitle', ''),
            'button_text': request.POST.get('button_text', ''),
            'button_url': request.POST.get('button_url', ''),
            'button2_text': request.POST.get('button2_text', ''),
            'button2_url': request.POST.get('button2_url', ''),
            'order': int(request.POST.get('order', 0)),
            'is_active': 'is_active' in request.POST,
        }
        if banner:
            for k, v in data.items():
                setattr(banner, k, v)
            if request.FILES.get('image'):
                banner.image = request.FILES['image']
            banner.save()
        else:
            banner = HeroBanner(**data)
            if request.FILES.get('image'):
                banner.image = request.FILES['image']
            banner.save()
        django_messages.success(request, 'Bannière sauvegardée !')
        return redirect('dashboard:banners')
    return render(request, 'dashboard/banner_edit.html', {'banner': banner})


@login_required
def admin_banner_delete(request, pk):
    if not (request.user.is_superuser or request.user.role == 'admin'):
        return redirect('dashboard:index')
    from catalog.models import HeroBanner
    banner = get_object_or_404(HeroBanner, pk=pk)
    if request.method == 'POST':
        banner.delete()
        django_messages.success(request, 'Bannière supprimée.')
    return redirect('dashboard:banners')

@login_required
def categories(request):
    """Gestion des catégories"""
    if not (request.user.is_superuser or request.user.role == 'admin'):
        return redirect('dashboard:index')

    categories = Category.objects.all().order_by('order', 'name')
    return render(request, 'dashboard/categories.html', {'categories': categories})

@login_required
def category_add(request):
    """Ajouter une catégorie"""
    if not (request.user.is_superuser or request.user.role == 'admin'):
        return redirect('dashboard:index')

    if request.method == 'POST':
        name = request.POST.get('name')
        icon = request.POST.get('icon', '📦')
        parent_id = request.POST.get('parent')
        order = request.POST.get('order', 0)
        description = request.POST.get('description', '')

        if name:
            slug = slugify(name)
            parent = Category.objects.get(pk=parent_id) if parent_id else None

            Category.objects.create(
                name=name,
                slug=slug,
                icon=icon,
                parent=parent,
                order=order,
                description=description
            )
            django_messages.success(request, f'Catégorie "{name}" créée avec succès!')
            return redirect('dashboard:categories')

    parent_categories = Category.objects.filter(parent__isnull=True)
    return render(request, 'dashboard/category_form.html', {
        'parent_categories': parent_categories,
        'title': 'Ajouter une catégorie'
    })

@login_required
def category_edit(request, pk):
    """Modifier une catégorie"""
    if not (request.user.is_superuser or request.user.role == 'admin'):
        return redirect('dashboard:index')

    category = get_object_or_404(Category, pk=pk)

    if request.method == 'POST':
        category.name = request.POST.get('name')
        category.slug = slugify(category.name)
        category.icon = request.POST.get('icon', '📦')
        category.description = request.POST.get('description', '')
        category.order = request.POST.get('order', 0)

        parent_id = request.POST.get('parent')
        category.parent = Category.objects.get(pk=parent_id) if parent_id else None

        category.save()
        django_messages.success(request, f'Catégorie "{category.name}" modifiée!')
        return redirect('dashboard:categories')

    parent_categories = Category.objects.filter(parent__isnull=True).exclude(pk=category.pk)
    return render(request, 'dashboard/category_form.html', {
        'category': category,
        'parent_categories': parent_categories,
        'title': 'Modifier la catégorie'
    })

@login_required
def category_delete(request, pk):
    """Supprimer une catégorie"""
    if not (request.user.is_superuser or request.user.role == 'admin'):
        return redirect('dashboard:index')

    category = get_object_or_404(Category, pk=pk)

    if request.method == 'POST':
        name = category.name
        category.delete()
        django_messages.success(request, f'Catégorie "{name}" supprimée!')
        return redirect('dashboard:categories')

    return render(request, 'dashboard/category_delete.html', {'category': category})
