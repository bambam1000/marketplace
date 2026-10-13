"""Import d'un catalogue fournisseur depuis un fichier CSV ou Excel.

1. read_file : le fichier devient une liste d'en-têtes et de lignes (enregistrées dans SupplierImport).
2. guess_mapping : chaque champ Comptoir est associé à une colonne d'après son titre.
3. run_import : les produits sont créés, ou mis à jour quand la référence fournisseur est déjà connue.
"""
import csv
import io
import re
import unicodedata
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.utils import timezone

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000

# (code, libellé, obligatoire, titres de colonnes reconnus)
FIELDS = [
    ('sku', 'Référence fournisseur', True, ['sku', 'reference', 'ref', 'code', 'code article', 'product id', 'id produit', 'article', 'id', 'variant sku']),
    ('name', 'Nom du produit', True, ['nom', 'name', 'titre', 'title', 'produit', 'product', 'product name', 'designation', 'libelle', 'nom du produit']),
    ('cost', 'Prix fournisseur', True, ['prix fournisseur', 'cout', 'cost', 'prix achat', 'prix d achat', 'purchase price', 'wholesale price', 'prix de gros', 'prix', 'price', 'unit price', 'sell price']),
    ('price', 'Prix de vente (facultatif)', False, ['prix de vente', 'prix vente', 'sale price', 'retail price', 'prix public', 'pvc', 'prix conseille']),
    ('stock', 'Stock', False, ['stock', 'quantite', 'qte', 'qty', 'quantity', 'inventory', 'inventaire', 'disponible']),
    ('description', 'Description', False, ['description', 'desc', 'details', 'detail', 'product description']),
    ('category', 'Catégorie', False, ['categorie', 'category', 'famille', 'rayon', 'type']),
    ('image', 'Image principale (lien)', False, ['image', 'image url', 'photo', 'img', 'main image', 'picture', 'image principale', 'url image']),
    ('images', 'Autres images (liens)', False, ['images', 'autres images', 'gallery', 'galerie', 'photos', 'additional images']),
    ('colors', 'Couleurs', False, ['couleur', 'couleurs', 'color', 'colors', 'colour']),
    ('sizes', 'Tailles', False, ['taille', 'tailles', 'size', 'sizes', 'pointure']),
    ('weight', 'Poids', False, ['poids', 'weight', 'poids kg', 'weight kg']),
    ('url', 'Lien chez le fournisseur', False, ['lien', 'url', 'link', 'product url', 'lien produit']),
]
FIELD_LABELS = {code: label for code, label, _, _ in FIELDS}
REQUIRED = [code for code, _, req, _ in FIELDS if req]


class ImportFileError(Exception):
    pass


def _norm(text):
    text = unicodedata.normalize('NFKD', str(text or '')).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'[^a-z0-9]+', ' ', text).strip()


def _cell(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def read_file(upload):
    """Retourne (en-têtes, lignes). Lève ImportFileError avec un message lisible."""
    name = (upload.name or '').lower()
    if upload.size > MAX_BYTES:
        raise ImportFileError('Le fichier dépasse 5 Mo. Découpez-le en plusieurs fichiers.')
    data = upload.read()
    if name.endswith('.xlsx'):
        rows = _read_xlsx(data)
    elif name.endswith('.csv') or name.endswith('.txt'):
        rows = _read_csv(data)
    elif name.endswith('.xls'):
        raise ImportFileError("L'ancien format Excel (.xls) n'est pas lu : ouvrez le fichier et enregistrez-le en .xlsx ou en .csv.")
    else:
        raise ImportFileError('Format non reconnu : envoyez un fichier .csv ou .xlsx.')
    rows = [[_cell(c) for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if len(rows) < 2:
        raise ImportFileError('Le fichier ne contient pas de produits : il faut une ligne de titres, puis une ligne par produit.')
    headers, body = rows[0], rows[1:]
    if len(body) > MAX_ROWS:
        raise ImportFileError(f'Le fichier contient {len(body)} produits : la limite est de {MAX_ROWS} par import.')
    width = len(headers)
    headers = [h or f'Colonne {i + 1}' for i, h in enumerate(headers)]
    body = [(r + [''] * width)[:width] for r in body]
    return headers, body


def _read_xlsx(data):
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise ImportFileError("Ce fichier Excel n'a pas pu être ouvert. Est-il protégé ou abîmé ?")
    ws = wb.worksheets[0]
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append(list(row))
        if len(rows) > MAX_ROWS + 50:
            break
    wb.close()
    return rows


def _read_csv(data):
    for enc in ('utf-8-sig', 'cp1252', 'latin-1'):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    sample = text[:5000]
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=',;\t|').delimiter
    except csv.Error:
        delimiter = ';' if sample.count(';') > sample.count(',') else ','
    return list(csv.reader(io.StringIO(text), delimiter=delimiter))


def guess_mapping(headers):
    """{champ: index de colonne} d'après les titres. Une colonne ne sert qu'à un champ."""
    normed = [_norm(h) for h in headers]
    # Titre identique à un nom connu, sinon titre qui commence ou finit par ce nom ; le nom le plus long l'emporte
    # (« Prix de vente (F) » va au prix de vente, pas au prix fournisseur)
    candidates = []
    for code, _, _, synonyms in FIELDS:
        for i, h in enumerate(normed):
            for syn in synonyms:
                if h == syn:
                    candidates.append((2, len(syn), code, i))
                elif h.startswith(syn + ' ') or h.endswith(' ' + syn):
                    candidates.append((1, len(syn), code, i))
    mapping, used = {}, set()
    for _, _, code, i in sorted(candidates, reverse=True):
        if code not in mapping and i not in used:
            mapping[code] = i
            used.add(i)
    return mapping


def parse_number(text):
    """« 1 250,50 », « $12.5 », « 1.234,56 € », « 12 USD » → Decimal. None si illisible."""
    s = re.sub(r'[^\d,.\-]', '', str(text or '').replace('\xa0', ''))
    if not s or s in '-.,':
        return None
    if ',' in s and '.' in s:
        if s.rfind(',') > s.rfind('.'):
            s = s.replace('.', '').replace(',', '.')
        else:
            s = s.replace(',', '')
    elif ',' in s:
        head, _, tail = s.rpartition(',')
        s = s.replace(',', '.') if s.count(',') == 1 and len(tail) <= 2 else s.replace(',', '')
    elif s.count('.') > 1:
        s = s.replace('.', '')
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def _split_list(text):
    return [p.strip() for p in re.split(r'[|;,\n]', text or '') if p.strip()]


def _get(row, mapping, code):
    idx = mapping.get(code)
    return row[idx].strip() if idx is not None and idx < len(row) else ''


def preview(imp, mapping, connection, limit=8):
    """Premières lignes telles qu'elles seront importées : [(n° de ligne, valeurs, erreur)]."""
    out = []
    for n, row in enumerate(imp.rows[:limit], start=2):
        values, error = _row_values(row, mapping, connection)
        out.append((n, values, error))
    return out


def _row_values(row, mapping, connection):
    """Valeurs d'une ligne et prix calculés. Retourne (valeurs, erreur ou '')."""
    v = {code: _get(row, mapping, code) for code in FIELD_LABELS}
    if not v['sku']:
        return v, 'Référence fournisseur vide'
    if not v['name']:
        return v, 'Nom vide'
    cost = parse_number(v['cost'])
    if cost is None or cost < 0:
        return v, f'Prix fournisseur illisible : « {v["cost"]} »' if v['cost'] else 'Prix fournisseur vide'
    v['cost_value'] = cost
    v['landed'] = connection.landed_cost(cost)
    fixed = parse_number(v['price']) if v['price'] else None
    v['fixed_price'] = fixed is not None and fixed > 0
    v['sale'] = int(fixed) if v['fixed_price'] else connection.sale_price(cost)
    v['margin'] = connection.margin_pct(v['sale'], v['landed'])
    stock = parse_number(v['stock']) if v['stock'] else None
    v['stock_value'] = max(0, int(stock)) if stock is not None else None
    return v, ''


def _match_category(name, categories, default):
    if name:
        key = _norm(name.split('>')[-1])
        if key in categories:
            return categories[key], True
    return default, not name


def run_import(imp, mapping, options, user):
    """Crée ou met à jour les produits. Retourne l'import mis à jour."""
    from catalog.models import Category, Product
    from .models import SupplierProduct

    connection = imp.connection
    warehouse = connection.warehouse
    store = warehouse.store
    default_category = options['default_category']
    categories = {_norm(c.name): c for c in Category.objects.all()}
    default_stock = options['default_stock']
    report, created, updated, skipped = [], 0, 0, 0
    seen = set()
    now = timezone.now()
    existing = {sp.supplier_sku: sp for sp in connection.items.select_related('product')}
    # Produits en ligne permis par la formule : au-delà, les nouveaux produits sont importés hors ligne
    from billing.limits import limit_of
    used, maximum = limit_of(store, 'products')
    room = None if maximum is None else max(0, maximum - used)
    held_offline = 0

    for n, row in enumerate(imp.rows, start=2):
        v, error = _row_values(row, mapping, connection)
        if error:
            report.append([n, error])
            skipped += 1
            continue
        sku = v['sku'][:120]
        if sku in seen:
            report.append([n, f'Référence « {sku} » déjà présente plus haut dans le fichier : ligne ignorée'])
            skipped += 1
            continue
        seen.add(sku)
        stock = v['stock_value'] if v['stock_value'] is not None else default_stock
        images = [u for u in [v['image']] + _split_list(v['images']) if u.lower().startswith(('http://', 'https://'))][:4]
        if v['fixed_price'] and v['margin'] < connection.min_margin_pct:
            report.append([n, f'« {v["name"][:60]} » : marge de {v["margin"]} %, sous votre seuil de {connection.min_margin_pct} %'])
        with transaction.atomic():
            sp = existing.get(sku)
            if sp:
                if not options['update_existing']:
                    skipped += 1
                    continue
                product = sp.product
                sp.cost, sp.supplier_stock, sp.last_synced_at = v['cost_value'], stock, now
                if v['url']:
                    sp.supplier_url = v['url'][:1000]
                sp.save()
                fields = ['cost_price']
                product.cost_price = v['landed']
                if sp.auto_price or v['fixed_price']:
                    product.price = v['sale']
                    fields.append('price')
                product.save(update_fields=fields)
                delta = stock - product.stock
                if delta:
                    product.adjust_stock(delta, 'adjustment', user=user, reason='Mise à jour du stock fournisseur',
                                         reference=connection.name[:50], warehouse=warehouse)
                updated += 1
            else:
                category, known = _match_category(v['category'], categories, default_category)
                if not known:
                    report.append([n, f'Catégorie « {v["category"]} » inconnue : produit rangé dans « {default_category.name} »'])
                publish = options['publish']
                if publish and room is not None:
                    if room > 0:
                        room -= 1
                    else:
                        publish = False
                        held_offline += 1
                product = Product.objects.create(
                    store=store, category=category, name=v['name'][:500],
                    description=v['description'] or v['name'],
                    price=v['sale'], cost_price=v['landed'], stock=0, sku=sku[:50],
                    colors=', '.join(_split_list(v['colors']))[:300],
                    sizes=', '.join(_split_list(v['sizes']))[:300],
                    weight=v['weight'][:50], origin=connection.country[:100] or 'Importé',
                    is_active=publish,
                )
                SupplierProduct.objects.create(
                    connection=connection, product=product, supplier_sku=sku, supplier_url=v['url'][:1000],
                    cost=v['cost_value'], supplier_stock=stock, auto_price=not v['fixed_price'],
                    image_urls=images, last_synced_at=now,
                )
                if stock:
                    product.adjust_stock(stock, 'in', user=user, reason='Import du catalogue fournisseur',
                                         reference=connection.name[:50], warehouse=warehouse)
                else:
                    from inventory.models import ProductStock
                    ProductStock.objects.get_or_create(product=product, warehouse=warehouse)
                created += 1

    if held_offline:
        report.append([0, f'{held_offline} nouveau{"x" if held_offline > 1 else ""} produit{"s" if held_offline > 1 else ""} '
                          f'importé{"s" if held_offline > 1 else ""} hors ligne : votre formule limite le nombre de produits en ligne '
                          f'({maximum}). Passez à une formule supérieure pour les publier.'])

    zeroed = 0
    if options['zero_missing']:
        for sku, sp in existing.items():
            if sku not in seen and sp.product.stock:
                sp.product.adjust_stock(-sp.product.stock, 'adjustment', user=user,
                                        reason='Absent du dernier fichier fournisseur',
                                        reference=connection.name[:50], warehouse=warehouse)
                sp.supplier_stock = 0
                sp.save(update_fields=['supplier_stock'])
                zeroed += 1
        if zeroed:
            report.append([0, f'{zeroed} produit{"s" if zeroed > 1 else ""} absent{"s" if zeroed > 1 else ""} du fichier : stock mis à zéro'])

    imp.mapping = {k: v for k, v in mapping.items()}
    imp.options = {'default_category': default_category.pk, 'default_stock': default_stock,
                   'update_existing': options['update_existing'], 'zero_missing': options['zero_missing'],
                   'publish': options['publish']}
    imp.created_count, imp.updated_count, imp.skipped_count = created, updated, skipped
    imp.report = report[:500]
    imp.status = 'done'
    imp.finished_at = timezone.now()
    imp.rows = []  # le fichier n'est plus utile : on ne garde que le compte rendu
    imp.save()
    connection.last_import_at = imp.finished_at
    connection.save(update_fields=['last_import_at'])
    return imp


def template_workbook():
    """Fichier modèle à remplir par le vendeur ou à envoyer au grossiste."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = 'Produits'
    heads = ['Référence', 'Nom', 'Prix fournisseur', 'Prix de vente', 'Stock', 'Description', 'Catégorie',
             'Image', 'Autres images', 'Couleurs', 'Tailles', 'Poids', 'Lien']
    ws.append(heads)
    ws.append(['CJ-1001', 'Montre connectée sport', '12.50', '', '150', 'Écran tactile, autonomie 7 jours',
               'Électronique', 'https://exemple.com/montre.jpg', 'https://exemple.com/montre-2.jpg',
               'Noir, Bleu', '', '0.2 kg', 'https://exemple.com/produit/1001'])
    ws.append(['CJ-1002', 'Sac à dos imperméable', '8', '9900', '80', '', 'Mode', '', '', 'Gris', '', '0.6 kg', ''])
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='FF6A00')
    for col, width in zip('ABCDEFGHIJKLM', (14, 30, 16, 14, 8, 36, 16, 34, 34, 14, 12, 10, 34)):
        ws.column_dimensions[col].width = width
    notes = wb.create_sheet('Mode d\'emploi')
    for line in [
        'Une ligne par produit. Seules les colonnes Référence, Nom et Prix fournisseur sont obligatoires.',
        "Prix fournisseur : dans la devise choisie dans les réglages du fournisseur (ex. 12.50 pour 12,50 $).",
        'Prix de vente : facultatif. Laissé vide, il est calculé avec votre coefficient et vos frais.',
        'Stock : facultatif. Laissé vide, le stock par défaut choisi au moment de l\'import est utilisé.',
        'Images : liens commençant par http:// ou https://. Plusieurs liens séparés par une virgule.',
        "Réimporter le même fichier met à jour les prix et le stock des produits déjà importés (reconnus par la référence).",
    ]:
        notes.append([line])
    notes.column_dimensions['A'].width = 110
    return wb
