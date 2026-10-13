"""Pages « Fournisseur » d'un entrepôt en dropshipping (dashboard vendeur)."""
from datetime import timedelta
from decimal import Decimal, InvalidOperation

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Count, F, Q, Sum
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.http import require_POST

from catalog.models import Category
from dashboard.views import seller_or_admin_required
from store import access

from . import importer
from .images import fetch_pending
from .models import SupplierConnection, SupplierImport, SupplierOrder, SupplierProduct
from .countries import country_name
from .services import mark_order_shipped, notify_buyer_shipped

IMAGE_BUDGET = 25  # secondes de téléchargement d'images pendant une requête


def _connection_or_deny(request, pk, permission):
    """(connexion, None) si l'utilisateur a le droit, sinon (None, redirection)."""
    from inventory.models import Warehouse
    warehouse = get_object_or_404(Warehouse.objects.select_related('store'), pk=pk)
    if not warehouse.is_dropship:
        return None, redirect('dashboard:warehouse_detail', pk=pk)
    if not access.stock_access(request.user, warehouse.store, permission, warehouse):
        messages.error(request, 'Accès refusé.')
        return None, redirect('dashboard:warehouses')
    conn, _ = SupplierConnection.objects.get_or_create(warehouse=warehouse, defaults={'name': warehouse.name})
    return conn, None


def _ctx(request, conn, tab, **extra):
    perm = lambda p: access.stock_access(request.user, conn.warehouse.store, p, conn.warehouse)  # noqa: E731
    paid_to_send = conn.orders.filter(status='to_send', order__is_paid=True).exclude(order__status__in=('cancelled', 'refunded')).count()
    return {
        'conn': conn, 'warehouse': conn.warehouse, 'tab': tab,
        'can': {'settings': perm(None), 'import': perm('products.edit'), 'orders': perm('orders.view'),
                'manage_orders': perm('orders.manage')},
        'badges': {'to_send': paid_to_send, 'to_cancel': conn.orders.filter(status='to_cancel').count()},
        **extra,
    }


# ─── Accueil : liste des fournisseurs ───

@login_required
@seller_or_admin_required
def index(request):
    """Tous les fournisseurs de la boutique ; création rapide d'une connexion CJ ou d'un fournisseur par fichier."""
    from inventory.models import Warehouse
    store = access.acting_store(request.user)
    if store is None:
        return redirect('dashboard:warehouses')
    owner = access.has_perm(request.user, store, None)
    if request.method == 'POST':
        if not owner:
            messages.error(request, 'Seul le propriétaire peut ajouter un fournisseur.')
            return redirect('dropship:index')
        provider = 'cj' if request.POST.get('provider') == 'cj' else 'file'
        name = (request.POST.get('name') or '').strip()[:100] or ('CJ Dropshipping' if provider == 'cj' else '')
        if len(name) < 2:
            messages.error(request, 'Indiquez le nom du fournisseur.')
            return redirect('dropship:index')
        base = 'CJ' if provider == 'cj' else 'FRN'
        n = 1
        while store.warehouses.filter(code=f'{base}-{n:02d}').exists():
            n += 1
        wh = Warehouse.objects.create(store=store, kind='dropship', name=name, code=f'{base}-{n:02d}',
                                      city=store.city or 'Douala')
        SupplierConnection.objects.create(warehouse=wh, name=name, provider=provider, currency='USD')
        messages.success(request, f'Fournisseur « {name} » créé.' + (" Collez votre clé d'API CJ pour le connecter." if provider == 'cj' else ''))
        return redirect('dropship:settings', pk=wh.pk)
    whs = Warehouse.objects.filter(store=store, kind='dropship').select_related('supplier')
    limit = access.member_warehouse_id(request.user, store)
    if limit:
        whs = whs.filter(pk=limit)
    suppliers = []
    for wh in whs:
        conn = getattr(wh, 'supplier', None) or SupplierConnection.objects.create(warehouse=wh, name=wh.name)
        suppliers.append({
            'conn': conn, 'wh': wh, 'products': conn.items.count(),
            'to_send': conn.orders.filter(status='to_send', order__is_paid=True).exclude(order__status__in=('cancelled', 'refunded')).count(),
            'transit': conn.orders.filter(status__in=('sent', 'shipped')).count(),
            'to_cancel': conn.orders.filter(status='to_cancel').count(),
        })
    return render(request, 'dropship/index.html', {'suppliers': suppliers, 'owner': owner,
                                                    'has_cj': any(x['conn'].provider == 'cj' for x in suppliers)})


@login_required
@seller_or_admin_required
def guide(request):
    """Mode d'emploi : CJ Dropshipping, AliExpress et fournisseurs par fichier."""
    from inventory.models import Warehouse
    store = access.acting_store(request.user)
    suppliers = Warehouse.objects.filter(store=store, kind='dropship').select_related('supplier') if store else []
    first = {'cj': None, 'file': None}
    for wh in suppliers:
        kind = 'cj' if getattr(wh, 'supplier', None) and wh.supplier.provider == 'cj' else 'file'
        first[kind] = first[kind] or wh
    return render(request, 'dropship/guide.html', {'cj_wh': first['cj'], 'file_wh': first['file'],
                                                    'owner': store is not None and access.has_perm(request.user, store, None)})


# ─── Réglages ───

SETTINGS_NUMBERS = [
    # (champ, libellé, min, max, décimal)
    ('exchange_rate', 'Taux de change', Decimal('0.0001'), Decimal('100000'), True),
    ('fixed_fee', 'Frais par article', 0, 10_000_000, False),
    ('markup', 'Coefficient', Decimal('1'), Decimal('20'), True),
    ('round_to', 'Arrondi', 1, 100_000, False),
    ('min_margin_pct', 'Marge minimale', 0, 95, False),
    ('delivery_min_days', 'Délai minimum', 1, 120, False),
    ('delivery_max_days', 'Délai maximum', 1, 180, False),
]


@login_required
@seller_or_admin_required
def settings_view(request, pk):
    conn, denied = _connection_or_deny(request, pk, None)
    if denied:
        return denied
    errors = {}
    if request.method == 'POST':
        post = request.POST
        if post.get('action') in ('cj_save', 'cj_disconnect'):
            return _cj_settings(request, conn)
        if post.get('action') == 'reprice':
            n = _reprice(conn.items.filter(auto_price=True).select_related('product'), conn)
            messages.success(request, f'{n} prix de vente recalculé{"s" if n > 1 else ""} avec les réglages actuels.')
            return redirect('dropship:settings', pk=pk)
        name = (post.get('name') or '').strip()
        if len(name) < 2:
            errors['name'] = 'Indiquez le nom du fournisseur.'
        provider = post.get('provider') if post.get('provider') in dict(SupplierConnection.PROVIDER_CHOICES) else 'file'
        currency = post.get('currency') if post.get('currency') in dict(SupplierConnection.CURRENCY_CHOICES) else 'USD'
        values = {}
        for field, label, lo, hi, dec in SETTINGS_NUMBERS:
            raw = (post.get(field) or '').replace(' ', '').replace(',', '.')
            if field == 'exchange_rate' and currency == 'XAF':
                values[field] = Decimal(1)
                continue
            try:
                val = Decimal(raw) if dec else int(Decimal(raw))
            except (InvalidOperation, ValueError):
                errors[field] = f'{label} : indiquez un nombre.'
                continue
            if not lo <= val <= hi:
                errors[field] = f'{label} : entre {lo} et {hi}.'
            values[field] = val
        if not errors and values['delivery_min_days'] > values['delivery_max_days']:
            errors['delivery_max_days'] = 'Le délai maximum doit être supérieur au délai minimum.'
        if not errors:
            pricing_changed = any(getattr(conn, f) != values[f] for f in ('exchange_rate', 'fixed_fee', 'markup', 'round_to')) \
                or conn.currency != currency
            conn.name, conn.provider, conn.currency = name[:150], provider, currency
            conn.contact_email = (post.get('contact_email') or '').strip()[:254]
            conn.contact_phone = (post.get('contact_phone') or '').strip()[:30]
            conn.country = (post.get('country') or '').strip()[:100] or 'Chine'
            conn.notes = (post.get('notes') or '').strip()
            for f, v in values.items():
                setattr(conn, f, v)
            conn.save()
            wh = conn.warehouse
            if wh.name != conn.name and post.get('rename_warehouse'):
                wh.name = conn.name[:100]
                wh.save(update_fields=['name'])
            msg = 'Réglages du fournisseur enregistrés.'
            if pricing_changed and conn.items.filter(auto_price=True).exists():
                msg += ' Les prix déjà en ligne ne changent pas : utilisez « Recalculer les prix » pour les appliquer.'
            messages.success(request, msg)
            return redirect('dropship:settings', pk=pk)
        conn = SupplierConnection(pk=conn.pk, warehouse=conn.warehouse, **{
            'name': name, 'provider': provider, 'currency': currency,
            'contact_email': post.get('contact_email', ''), 'contact_phone': post.get('contact_phone', ''),
            'country': post.get('country', ''), 'notes': post.get('notes', ''),
        })
        for field, *_ in SETTINGS_NUMBERS:
            setattr(conn, field, post.get(field, ''))
    example = None
    if not errors:
        cost = Decimal(10) if conn.currency != 'XAF' else Decimal(5000)
        example = {'cost': cost, 'landed': conn.landed_cost(cost), 'sale': conn.sale_price(cost)}
        example['margin'] = conn.margin_pct(example['sale'], example['landed'])
    return render(request, 'dropship/settings.html', _ctx(
        request, conn, 'settings', errors=errors, example=example,
        providers=SupplierConnection.PROVIDER_CHOICES, currencies=SupplierConnection.CURRENCY_CHOICES,
        countries=_country_choices(), frequent_codes=_country_choices()['codes'], ship_country=country_name(conn.ship_to_country),
        auto_count=conn.items.filter(auto_price=True).count() if conn.pk else 0,
    ), status=400 if errors else 200)


def _country_choices():
    from .countries import COUNTRIES, FREQUENT, NAMES
    return {'frequent': [(c, NAMES[c]) for c in FREQUENT], 'all': COUNTRIES, 'codes': FREQUENT}


def _cj_settings(request, conn):
    """Formulaire « Connexion directe CJ » des réglages."""
    from . import secret
    from .cj import CJClient, CJError
    post = request.POST
    if post.get('action') == 'cj_disconnect':
        conn.api_key_enc = conn.access_token_enc = conn.refresh_token_enc = conn.api_key_hint = ''
        conn.token_expires_at = conn.refresh_expires_at = conn.api_connected_at = None
        conn.auto_send = False
        conn.save()
        messages.success(request, 'Connexion CJ supprimée : la clé et les jetons ont été effacés.')
        return redirect('dropship:settings', pk=conn.warehouse_id)
    conn.logistic_name = (post.get('logistic_name') or '').strip()[:100] or 'CJPacket Ordinary'
    from .countries import NAMES
    country = (post.get('ship_to_country') or '').strip().upper()
    conn.ship_to_country = country if country in NAMES else 'CM'
    conn.include_shipping = 'include_shipping' in post
    conn.pay_from_balance = 'pay_from_balance' in post
    key = (post.get('api_key') or '').strip()
    if key:
        conn.provider = 'cj'
        try:
            CJClient(conn).authenticate(api_key=key)
        except CJError as exc:
            conn.refresh_from_db()
            messages.error(request, f"Connexion refusée : {exc}")
            return redirect('dropship:settings', pk=conn.warehouse_id)
        conn.api_key_enc, conn.api_key_hint = secret.encrypt(key), key[-4:]
        conn.api_connected_at, conn.api_error = timezone.now(), ''
        if conn.currency != 'USD':
            conn.currency = 'USD'
            messages.info(request, 'Devise passée en dollars : CJ donne ses prix en USD. Vérifiez le taux de change.')
    elif not conn.api_key_enc:
        messages.error(request, "Collez votre clé d'API CJ (compte CJ > Authorization > API).")
        return redirect('dropship:settings', pk=conn.warehouse_id)
    conn.auto_send = 'auto_send' in post
    if conn.auto_send:
        from billing.limits import blocked
        reason = blocked(conn.warehouse.store, feature='dropship_auto', user=request.user)
        if reason:
            conn.auto_send = False
            messages.warning(request, f'Transmission automatique non activée. {reason}')
    conn.save()
    messages.success(request, 'Connecté à CJ Dropshipping.' if key else 'Options de connexion enregistrées.')
    return redirect('dropship:settings', pk=conn.warehouse_id)


def _reprice(items, conn):
    n = 0
    for sp in items:
        new = conn.sale_price(sp.total_cost)
        landed = conn.landed_cost(sp.total_cost)
        p = sp.product
        if p.price != new or p.cost_price != landed:
            p.price, p.cost_price = new, landed
            p.save(update_fields=['price', 'cost_price'])
            n += 1
    return n


# ─── Import CSV / Excel ───

@login_required
@seller_or_admin_required
def import_upload(request, pk):
    conn, denied = _connection_or_deny(request, pk, 'products.edit')
    if denied:
        return denied
    error = ''
    if request.method == 'POST':
        upload = request.FILES.get('file')
        if not upload:
            error = 'Choisissez un fichier .csv ou .xlsx.'
        else:
            try:
                headers, rows = importer.read_file(upload)
            except importer.ImportFileError as exc:
                error = str(exc)
            else:
                imp = SupplierImport.objects.create(connection=conn, created_by=request.user, filename=upload.name[:255],
                                                    headers=headers, rows=rows, mapping=importer.guess_mapping(headers))
                return redirect('dropship:import_map', pk=pk, imp=imp.pk)
    # Imports abandonnés à l'étape des colonnes : on ne garde pas le contenu des fichiers indéfiniment
    conn.imports.filter(status='mapping', created_at__lt=timezone.now() - timedelta(days=2)).delete()
    return render(request, 'dropship/import_upload.html', _ctx(
        request, conn, 'import', error=error, imports=conn.imports.select_related('created_by')[:15],
        fields=importer.FIELDS,
    ), status=400 if error else 200)


def _read_mapping(post, width):
    mapping = {}
    for code in importer.FIELD_LABELS:
        raw = post.get(f'map_{code}', '')
        if raw.isdigit() and int(raw) < width:
            mapping[code] = int(raw)
    return mapping


@login_required
@seller_or_admin_required
def import_map(request, pk, imp):
    conn, denied = _connection_or_deny(request, pk, 'products.edit')
    if denied:
        return denied
    imp = get_object_or_404(SupplierImport, pk=imp, connection=conn)
    if imp.status != 'mapping':
        return redirect('dropship:import_result', pk=pk, imp=imp.pk)
    categories = Category.objects.filter(is_active=True).select_related('parent').order_by('parent__name', 'name')
    mapping = imp.mapping
    opts = {'default_category': '', 'default_stock': 50, 'update_existing': True, 'zero_missing': False, 'publish': True}
    errors = []
    if request.method == 'POST':
        mapping = _read_mapping(request.POST, len(imp.headers))
        post = request.POST
        opts = {
            'default_category': post.get('default_category', ''),
            'default_stock': post.get('default_stock', '0'),
            'update_existing': 'update_existing' in post, 'zero_missing': 'zero_missing' in post,
            'publish': 'publish' in post,
        }
        imp.mapping = mapping
        imp.save(update_fields=['mapping'])
        if post.get('action') == 'run':
            missing = [importer.FIELD_LABELS[c] for c in importer.REQUIRED if c not in mapping]
            if missing:
                errors.append('Associez une colonne à : ' + ', '.join(missing) + '.')
            used = list(mapping.values())
            if len(used) != len(set(used)):
                errors.append('Une même colonne est associée à deux champs.')
            category = categories.filter(pk=opts['default_category']).first() if str(opts['default_category']).isdigit() else None
            if not category:
                errors.append('Choisissez la catégorie par défaut.')
            try:
                stock = int(opts['default_stock'])
                if not 0 <= stock <= 1_000_000:
                    raise ValueError
            except (TypeError, ValueError):
                errors.append('Stock par défaut : indiquez un nombre entier positif.')
                stock = 0
            if not errors:
                options = {**opts, 'default_category': category, 'default_stock': stock}
                importer.run_import(imp, mapping, options, request.user)
                pending = conn.items.exclude(image_urls=[]).select_related('product')
                saved, left = fetch_pending(pending, budget_seconds=IMAGE_BUDGET)
                msg = f'Import terminé : {imp.created_count} produit{"s" if imp.created_count > 1 else ""} créé{"s" if imp.created_count > 1 else ""}, {imp.updated_count} mis à jour.'
                if saved:
                    msg += f' {saved} image{"s" if saved > 1 else ""} récupérée{"s" if saved > 1 else ""}.'
                messages.success(request, msg)
                return redirect('dropship:import_result', pk=pk, imp=imp.pk)
    already = set(conn.items.values_list('supplier_sku', flat=True))
    sku_col = mapping.get('sku')
    n_existing = sum(1 for r in imp.rows if sku_col is not None and r[sku_col].strip()[:120] in already)
    return render(request, 'dropship/import_map.html', _ctx(
        request, conn, 'import', imp=imp, mapping=mapping, opts=opts, errors=errors,
        field_rows=[(code, label, req, mapping.get(code)) for code, label, req, _ in importer.FIELDS], categories=categories, n_rows=len(imp.rows), n_existing=n_existing,
        preview=importer.preview(imp, mapping, conn) if all(c in mapping for c in importer.REQUIRED) else None,
        header_samples=[(i, h, next((r[i] for r in imp.rows if r[i]), '')[:40]) for i, h in enumerate(imp.headers)],
    ), status=400 if errors else 200)


@login_required
@seller_or_admin_required
def import_result(request, pk, imp):
    conn, denied = _connection_or_deny(request, pk, 'products.edit')
    if denied:
        return denied
    imp = get_object_or_404(SupplierImport, pk=imp, connection=conn)
    if imp.status == 'mapping':
        return redirect('dropship:import_map', pk=pk, imp=imp.pk)
    return render(request, 'dropship/import_result.html', _ctx(
        request, conn, 'import', imp=imp, images_left=conn.items.exclude(image_urls=[]).count(),
    ))


@login_required
@seller_or_admin_required
def import_template(request, pk):
    conn, denied = _connection_or_deny(request, pk, 'products.edit')
    if denied:
        return denied
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="modele_catalogue_fournisseur.xlsx"'
    importer.template_workbook().save(response)
    return response


@login_required
@seller_or_admin_required
@require_POST
def fetch_images(request, pk):
    conn, denied = _connection_or_deny(request, pk, 'products.edit')
    if denied:
        return denied
    saved, left = fetch_pending(conn.items.exclude(image_urls=[]).select_related('product'), budget_seconds=IMAGE_BUDGET)
    msg = f'{saved} image{"s" if saved > 1 else ""} récupérée{"s" if saved > 1 else ""}.'
    if left:
        msg += f' Encore {left} produit{"s" if left > 1 else ""} à traiter : relancez.'
    messages.success(request, msg)
    nxt = request.POST.get('next', '')
    return redirect(nxt if nxt.startswith('/') and not nxt.startswith('//') else f'/dashboard/dropshipping/{pk}/produits/')


# ─── Catalogue CJ (connexion directe) ───

CJ_IMPORT_MAX = 10


@login_required
@seller_or_admin_required
def cj_catalog(request, pk):
    from .cj import CJClient, CJError
    from .cj_sync import import_product
    conn, denied = _connection_or_deny(request, pk, 'products.edit')
    if denied:
        return denied
    if not conn.api_ready:
        messages.info(request, "Connectez votre compte CJ pour parcourir son catalogue.")
        return redirect('dropship:settings', pk=pk)
    categories = Category.objects.filter(is_active=True).select_related('parent').order_by('parent__name', 'name')
    if request.method == 'POST':
        pids = [p.strip()[:100] for p in request.POST.getlist('pids') if p.strip()]
        category = categories.filter(pk=request.POST.get('category')).first() if (request.POST.get('category') or '').isdigit() else None
        nxt = request.POST.get('next', '')
        back = redirect(nxt if nxt.startswith('/') and not nxt.startswith('//') else f'/dashboard/dropshipping/{pk}/catalogue/')
        if not pids:
            messages.error(request, 'Cochez au moins un produit à importer.')
            return back
        if not category:
            messages.error(request, 'Choisissez la catégorie Comptoir où ranger les produits.')
            return back
        if len(pids) > CJ_IMPORT_MAX:
            messages.error(request, f'Importez au plus {CJ_IMPORT_MAX} produits à la fois.')
            return back
        client = CJClient(conn)
        created = updated = 0
        for pid in pids:
            try:
                sp, new, warning = import_product(conn, pid, category, 'publish' in request.POST, request.user, client=client)
            except CJError as exc:
                messages.error(request, f'Produit {pid[:12]}… : {exc}')
                continue
            created += new
            updated += not new
            if warning:
                messages.warning(request, f'« {sp.product.name[:50]} » : port non calculé ({warning}).')
        saved, _ = fetch_pending(conn.items.exclude(image_urls=[]).select_related('product'), budget_seconds=15)
        if created or updated:
            conn.last_import_at = timezone.now()
            conn.save(update_fields=['last_import_at'])
            messages.success(request, f'{created} produit{"s" if created > 1 else ""} importé{"s" if created > 1 else ""}'
                             + (f', {updated} mis à jour' if updated else '') + (f', {saved} image{"s" if saved > 1 else ""}' if saved else '')
                             + '. Pensez à traduire les noms en français.')
        return back
    q = request.GET.get('q', '').strip()[:100]
    try:
        page_no = max(1, min(int(request.GET.get('page', 1)), 500))
    except ValueError:
        page_no = 1
    result, error = None, ''
    if q:
        try:
            result = CJClient(conn).search(q, page_no, 20)
        except CJError as exc:
            error = str(exc)
        else:
            known = dict(conn.items.filter(supplier_pid__in=[i['pid'] for i in result['items']]).values_list('supplier_pid', 'product_id'))
            for it in result['items']:
                it['estimate'] = conn.sale_price(it['cost'])
                it['known'] = known.get(it['pid'])
    return render(request, 'dropship/cj_catalog.html', _ctx(
        request, conn, 'catalog', q=q, result=result, error=error, page_no=page_no, categories=categories,
        max_import=CJ_IMPORT_MAX,
    ))


# ─── Produits ───

@login_required
@seller_or_admin_required
def products(request, pk):
    conn, denied = _connection_or_deny(request, pk, 'products.view')
    if denied:
        return denied
    can_edit = access.stock_access(request.user, conn.warehouse.store, 'products.edit', conn.warehouse)
    if request.method == 'POST':
        if not can_edit:
            messages.error(request, "Vous n'avez pas le droit de modifier les produits.")
            return redirect('dropship:products', pk=pk)
        ids = [int(i) for i in request.POST.getlist('ids') if i.isdigit()]
        chosen = conn.items.filter(pk__in=ids).select_related('product')
        action = request.POST.get('action')
        if request.POST.get('scope') == 'all':
            chosen = conn.items.order_by('last_synced_at').select_related('product')
        if not chosen:
            messages.error(request, 'Cochez au moins un produit.')
        elif action in ('publish', 'unpublish'):
            from catalog.models import Product
            n = Product.objects.filter(pk__in=chosen.values('product')).update(is_active=action == 'publish')
            messages.success(request, f'{n} produit{"s" if n > 1 else ""} {"publié" if action == "publish" else "retiré du site"}{"s" if n > 1 else ""}.')
        elif action == 'sync':
            from .cj_sync import sync_items
            if not conn.api_ready:
                messages.error(request, "Connectez d'abord CJ dans les réglages.")
            else:
                done, left, errors = sync_items(conn, list(chosen), request.user, budget_seconds=IMAGE_BUDGET)
                msg = f'{done} produit{"s" if done > 1 else ""} synchronisé{"s" if done > 1 else ""} avec CJ (prix, stock, variantes).'
                if left:
                    msg += f' Encore {left} à traiter : relancez.'
                messages.success(request, msg)
                for e in errors[:3]:
                    messages.error(request, e)
        elif action == 'reprice':
            n = _reprice(chosen, conn)
            chosen.update(auto_price=True)
            messages.success(request, f'Prix recalculés avec vos réglages ({n} modifié{"s" if n > 1 else ""}).')
        nxt = request.POST.get('next', '')
        return redirect(nxt if nxt.startswith('/') and not nxt.startswith('//') else f'/dashboard/dropshipping/{pk}/produits/')

    items = conn.items.select_related('product', 'product__category').order_by('product__name')
    min_margin = conn.min_margin_pct
    low_margin = Q(product__price__lt=F('product__cost_price') * 100 / (100 - min_margin)) if min_margin < 100 else Q(pk=None)
    counts = items.aggregate(
        all=Count('id'), out=Count('id', filter=Q(product__stock=0)), hidden=Count('id', filter=Q(product__is_active=False)),
        low=Count('id', filter=low_margin), images=Count('id', filter=~Q(image_urls=[])),
    )
    f = request.GET.get('f', '')
    filters = {'out': Q(product__stock=0), 'hidden': Q(product__is_active=False), 'low': low_margin, 'images': ~Q(image_urls=[])}
    if f in filters:
        items = items.filter(filters[f])
    else:
        f = ''
    q = request.GET.get('q', '').strip()
    if q:
        items = items.filter(Q(product__name__icontains=q) | Q(supplier_sku__icontains=q))
    page = Paginator(items, 30).get_page(request.GET.get('page'))
    for sp in page:
        sp.landed = sp.product.cost_price or conn.landed_cost(sp.total_cost)
        sp.margin = conn.margin_pct(sp.product.price, sp.landed)
        sp.low = sp.margin < min_margin
    tabs = [('', 'Tous', counts['all']), ('out', 'En rupture', counts['out']), ('low', 'Marge faible', counts['low']),
            ('hidden', 'Hors ligne', counts['hidden']), ('images', 'Images en attente', counts['images'])]
    return render(request, 'dropship/products.html', _ctx(
        request, conn, 'products', page=page, f=f, q=q, tabs=tabs, counts=counts, can_edit=can_edit,
    ))


# ─── Commandes fournisseur ───

ORDER_TABS = [
    ('', 'À traiter'), ('to_send', 'À transmettre'), ('sent', 'Transmises'), ('shipped', 'Expédiées'),
    ('delivered', 'Livrées'), ('to_cancel', 'À annuler'), ('cancelled', 'Annulées'), ('all', 'Toutes'),
]


@login_required
@seller_or_admin_required
def orders(request, pk):
    conn, denied = _connection_or_deny(request, pk, 'orders.view')
    if denied:
        return denied
    qs = conn.orders.select_related('order', 'order__buyer')
    counts = dict(qs.order_by().values_list('status').annotate(n=Count('id')))
    status = request.GET.get('status', '')
    if status == 'all':
        pass
    elif status in dict(SupplierOrder.STATUS_CHOICES):
        qs = qs.filter(status=status)
    else:
        status = ''
        qs = qs.filter(status__in=('to_send', 'sent', 'to_cancel'))
    q = request.GET.get('q', '').strip()
    if q:
        qs = qs.filter(Q(reference__icontains=q) | Q(order__order_number__icontains=q) | Q(order__shipping_name__icontains=q)
                       | Q(tracking_number__icontains=q) | Q(supplier_reference__icontains=q))
    page = Paginator(qs, 25).get_page(request.GET.get('page'))
    open_n = sum(counts.get(s, 0) for s in ('to_send', 'sent', 'to_cancel'))
    tabs = [(code, label, open_n if code == '' else sum(counts.values()) if code == 'all' else counts.get(code, 0))
            for code, label in ORDER_TABS]
    stats = conn.orders.exclude(status='cancelled').aggregate(cost=Sum('cost'), n=Count('id'))
    return render(request, 'dropship/orders.html', _ctx(
        request, conn, 'orders', page=page, tabs=tabs, status=status, q=q, stats=stats,
    ))


@login_required
@seller_or_admin_required
def order_detail(request, pk, so):
    conn, denied = _connection_or_deny(request, pk, 'orders.view')
    if denied:
        return denied
    so = get_object_or_404(SupplierOrder.objects.select_related('order', 'order__buyer'), pk=so, connection=conn)
    can_manage = access.stock_access(request.user, conn.warehouse.store, 'orders.manage', conn.warehouse)
    if request.method == 'POST':
        if not can_manage:
            messages.error(request, "Vous n'avez pas le droit de gérer les commandes.")
            return redirect('dropship:order_detail', pk=pk, so=so.pk)
        action = request.POST.get('action')
        val = lambda k, n=100: (request.POST.get(k) or '').strip()[:n]  # noqa: E731
        with transaction.atomic():
            so = SupplierOrder.objects.select_for_update().select_related('order').get(pk=so.pk)
            if action == 'sent':
                if not so.can_send:
                    messages.error(request, "Cette commande ne peut pas être transmise : le client doit d'abord l'avoir payée."
                                   if so.status == 'to_send' else 'Cette commande a déjà été traitée.')
                else:
                    so.status, so.sent_at = 'sent', timezone.now()
                    so.supplier_reference = val('supplier_reference')
                    cost = request.POST.get('cost', '').replace(' ', '')
                    if cost.isdigit():
                        so.cost = int(cost)
                    so.save()
                    messages.success(request, f'{so.reference} marquée comme transmise à {conn.name}.')
            elif action == 'shipped':
                if so.status not in ('sent', 'shipped'):
                    messages.error(request, "Transmettez d'abord la commande au fournisseur.")
                elif not val('tracking_number'):
                    messages.error(request, 'Indiquez le numéro de suivi donné par le fournisseur.')
                else:
                    first = so.status == 'sent'
                    so.status = 'shipped'
                    so.carrier, so.tracking_number = val('carrier'), val('tracking_number')
                    url = val('tracking_url', 500)
                    so.tracking_url = url if url.startswith(('http://', 'https://')) else ''
                    so.shipped_at = so.shipped_at or timezone.now()
                    so.save()
                    mark_order_shipped(so, request.user)
                    if first:
                        notify_buyer_shipped(so)
                    messages.success(request, 'Suivi enregistré.' + (' Le client a été prévenu.' if first else ''))
            elif action == 'delivered':
                if so.status != 'shipped':
                    messages.error(request, "Seule une commande expédiée peut être marquée livrée.")
                else:
                    so.status = 'delivered'
                    so.save(update_fields=['status', 'updated_at'])
                    messages.success(request, 'Commande fournisseur marquée livrée.')
            elif action == 'cancelled':
                if so.status not in ('to_send', 'to_cancel', 'sent'):
                    messages.error(request, "Cette commande ne peut plus être annulée ici.")
                else:
                    so.status = 'cancelled'
                    so.save(update_fields=['status', 'updated_at'])
                    messages.success(request, f'{so.reference} annulée.')
            elif action in ('cj_send', 'cj_refresh', 'cj_cancel'):
                from .cj import CJError
                from .cj_sync import cancel_at_supplier, refresh_order, send_order
                try:
                    if action == 'cj_send':
                        send_order(so, request.user)
                        messages.success(request, f'Commande transmise à CJ (n° {so.supplier_reference}).'
                                         + ('' if conn.pay_from_balance else " Payez-la dans votre espace CJ pour qu'elle parte."))
                    elif action == 'cj_refresh':
                        changed = refresh_order(so, request.user)
                        messages.success(request, 'Suivi mis à jour : ' + (so.get_status_display().lower() if changed else 'pas de changement') + '.')
                    elif so.status == 'to_cancel' and so.supplier_reference:
                        cancel_at_supplier(so)
                        messages.success(request, 'Commande supprimée chez CJ.')
                except CJError as exc:
                    messages.error(request, str(exc))
            elif action == 'notes':
                so.notes = (request.POST.get('notes') or '').strip()[:2000]
                so.save(update_fields=['notes', 'updated_at'])
                messages.success(request, 'Note enregistrée.')
        return redirect('dropship:order_detail', pk=pk, so=so.pk)

    lines = list(so.lines)
    sold = sum(int(l.price * l.quantity) for l in lines)
    others = so.order.items.exclude(warehouse=conn.warehouse).count()
    return render(request, 'dropship/order_detail.html', _ctx(
        request, conn, 'orders', so=so, order=so.order, lines=lines, sold=sold, margin=sold - int(so.cost),
        others=others, can_manage=can_manage,
    ))


@login_required
@seller_or_admin_required
def order_sheet(request, pk, so):
    """Bon de commande Excel à envoyer au fournisseur."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    conn, denied = _connection_or_deny(request, pk, 'orders.view')
    if denied:
        return denied
    so = get_object_or_404(SupplierOrder.objects.select_related('order'), pk=so, connection=conn)
    order = so.order
    wb = Workbook()
    ws = wb.active
    ws.title = 'Bon de commande'
    bold = Font(bold=True)
    ws.append([f'Bon de commande {so.reference}'])
    ws['A1'].font = Font(bold=True, size=14)
    ws.append([f'Boutique : {conn.warehouse.store.name}', '', f'Date : {timezone.localdate():%d/%m/%Y}'])
    ws.append([])
    ws.append(['Livrer à'])
    ws['A4'].font = bold
    for label, value in (('Nom', order.shipping_name), ('Téléphone', order.shipping_phone),
                         ('Adresse', order.shipping_address), ('Ville', order.shipping_city), ('Pays', 'Cameroun')):
        ws.append([label, value])
    ws.append([])
    ws.append(['Référence fournisseur', 'Produit', 'Couleur', 'Taille', 'Quantité', 'Lien'])
    for cell in ws[ws.max_row]:
        cell.font = bold
    refs = {sp.product_id: sp for sp in conn.items.filter(product__in=[l.product_id for l in so.lines])}
    for l in so.lines:
        sp = refs.get(l.product_id)
        ws.append([sp.supplier_sku if sp else l.product.sku, l.product.name, l.color, l.size, l.quantity,
                   sp.supplier_url if sp else ''])
    for col, width in zip('ABCDEF', (22, 40, 12, 10, 10, 40)):
        ws.column_dimensions[col].width = width
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="bon_{so.reference}.xlsx"'
    wb.save(response)
    return response
